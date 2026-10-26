"""
AAS Validator Auto Fixer (v3) - core
====================================

추진단 Validator(IDTA 2-x/3-x/4-x + KOSMO-AAS/Submodel/SME/CD) 기준으로
반복되는 오류를 1차 자동 보정한다. UI 는 app.py (Streamlit).

v3 핵심 변경 (v2 대비)
---------------------
* IDTA 템플릿 / 표준사전(IRDI·admin-shell IRI) 식별자를 보호한다.
  - v2 는 모든 CD id 와 Submodel semanticId 를 smart-factory 형식으로
    덮어써서 DN/HD/TD 의 IDTA semanticId 가 깨졌다.
* 표준 semanticId 에 대응 CD 가 없을 때
  - IRDI  -> 같은 id 로 CD 생성 (예: Markings 0112/2///61360_7#AAS006#001)
  - IRI   -> smart-factory CD 생성 + 원래 IRI 를 supplementalSemanticId 와
             CD isCaseOf 로 유지 (예: AddressInformation)
* AAS 3.1 에서는 AASd-120 이 없어졌으므로 SML 항목 idShort 를 지우지 않는다
  (3.0 네임스페이스일 때만 제거).
* IDTA 2-14 대상 확대: 빈 assetType, 빈 File value 등.
* assetKind=Type, Submodel kind=Template, globalAssetId 형식(KOSMO-AAS #5/#6).
* 끊어진 ModelReference(존재하지 않는 Submodel id) 자동 복구.
* CD 내용(KOSMO-CD #3): description <-> definition 동기화, displayName 보충,
  요소 description = CD definition, 한글 사전(CSV) 반영, 한글 누락 경고.
* 정수 valueType 은 xs:integer, 실수는 xs:float 로 통일.
* AASX 패키지: 디렉터리 엔트리 제거, 참조 파일/relationship/Content_Types
  정합, 썸네일(KOSMO-AAS #1) 점검, 파일명 언더바 제거(선택).
* Validator 항목 번호 기준 점검 리포트.
"""

import copy
import io
import re
import zipfile
import xml.etree.ElementTree as ET



# =========================================================
# 기본 설정
# =========================================================

BASE_URL = "https://www.smart-factory.kr/ids"

DEFAULT_VERSION = "1"
DEFAULT_REVISION = "0"

IEC61360_TEMPLATE_ID = (
    "https://admin-shell.io/DataSpecificationTemplates/"
    "DataSpecificationIec61360/3/0"
)

# 자동으로 채운 텍스트 표시 (그대로 제출되지 않도록 눈에 띄게)
TBD_EN = "[TBD]"
TBD_KO = "[작성 필요]"

SAMPLE_DATETIME = "2026-09-08T09:00:00+09:00"
SAMPLE_DATE = "2026-09-08"
SAMPLE_TIME = "09:00:00"




NS_AAS30 = "https://admin-shell.io/aas/3/0"
NS_AAS31 = "https://admin-shell.io/aas/3/1"

# 표준(템플릿/사전) 식별자: 덮어쓰지 않는다
STANDARD_ID_PATTERNS = [
    re.compile(r"^\d{4}-\d#\d{2}-[A-Z]{3}\d{3}#\d{3}"),   # ECLASS IRDI
    re.compile(r"^0112/2///"),                              # IEC CDD IRDI
    re.compile(r"^https?://admin-shell\.io/", re.I),
    re.compile(r"^https?://api\.eclass-cdp\.com/", re.I),
    re.compile(r"^https?://(www\.)?eclass\.eu/", re.I),
    re.compile(r"^urn:", re.I),
]

IRDI_PATTERN = re.compile(r"^(\d{4}-\d#|0112/2///)")


def is_self_id(value):
    return bool(value) and value.startswith(BASE_URL)


def is_standard_id(value):
    if not value or is_self_id(value):
        return False
    return any(p.match(value) for p in STANDARD_ID_PATTERNS)


def is_irdi(value):
    return bool(value) and IRDI_PATTERN.match(value) is not None


CD_NAME_ALIAS = {
    "WorkorderID": "WorkOrderID",
    "Workorder": "WorkOrder",
    "WorkorderNumber": "ProductionOrderNumber",
    "ProductionorderNumber": "ProductionOrderNumber",
}


# =========================================================
# XSD 자식 요소 순서
# =========================================================

_REFERABLE_HEAD = [
    "extensions", "category", "idShort", "displayName", "description",
]

_SME_HEAD = _REFERABLE_HEAD + [
    "semanticId", "supplementalSemanticIds",
    "qualifiers", "embeddedDataSpecifications",
]

_REF_ORDER = ["type", "referredSemanticId", "keys"]


CHILD_ORDER = {
    "environment": [
        "assetAdministrationShells", "submodels", "conceptDescriptions",
    ],
    "assetAdministrationShell": _REFERABLE_HEAD + [
        "administration", "id", "embeddedDataSpecifications",
        "derivedFrom", "assetInformation", "submodels",
    ],
    "assetInformation": [
        "assetKind", "globalAssetId", "specificAssetIds",
        "assetType", "defaultThumbnail",
    ],
    "submodel": _REFERABLE_HEAD + [
        "administration", "id", "kind", "semanticId",
        "supplementalSemanticIds", "qualifiers",
        "embeddedDataSpecifications", "submodelElements",
    ],
    "conceptDescription": _REFERABLE_HEAD + [
        "administration", "id", "embeddedDataSpecifications", "isCaseOf",
    ],
    "administration": [
        "embeddedDataSpecifications", "version", "revision",
        "creator", "templateId",
    ],
    "reference": _REF_ORDER,
    "semanticId": _REF_ORDER,
    "semanticIdListElement": _REF_ORDER,
    "valueId": _REF_ORDER,
    "unitId": _REF_ORDER,
    "dataSpecification": _REF_ORDER,
    "globalAssetId": _REF_ORDER,
    "first": _REF_ORDER,
    "externalSubjectId": _REF_ORDER,
    "derivedFrom": _REF_ORDER,
    "messageBroker": _REF_ORDER,
    "creator": _REF_ORDER,
    "refersTo": _REF_ORDER,
    "second": _REF_ORDER,
    "observed": _REF_ORDER,
    "key": ["type", "value"],
    "property": _SME_HEAD + ["valueType", "value", "valueId"],
    "multiLanguageProperty": _SME_HEAD + ["value", "valueId"],
    "range": _SME_HEAD + ["valueType", "min", "max"],
    "file": _SME_HEAD + ["value", "contentType"],
    "blob": _SME_HEAD + ["value", "contentType"],
    "referenceElement": _SME_HEAD + ["value"],
    "relationshipElement": _SME_HEAD + ["first", "second"],
    "annotatedRelationshipElement": _SME_HEAD + [
        "first", "second", "annotations",
    ],
    "entity": _SME_HEAD + [
        "statements", "entityType", "globalAssetId", "specificAssetIds",
    ],
    "operation": _SME_HEAD + [
        "inputVariables", "outputVariables", "inoutputVariables",
    ],
    "capability": list(_SME_HEAD),
    "basicEventElement": _SME_HEAD + [
        "observed", "direction", "state", "messageTopic", "messageBroker",
        "lastUpdate", "minInterval", "maxInterval",
    ],
    "submodelElementCollection": _SME_HEAD + ["value"],
    "submodelElementList": _SME_HEAD + [
        "orderRelevant", "semanticIdListElement", "typeValueListElement",
        "valueTypeListElement", "value",
    ],
    # 실제 AASX 산출물 확인 결과 dataSpecification 이 먼저다.
    "embeddedDataSpecification": [
        "dataSpecification", "dataSpecificationContent",
    ],
    "dataSpecificationIec61360": [
        "preferredName", "shortName", "unit", "unitId", "sourceOfDefinition",
        "symbol", "dataType", "definition", "valueFormat", "valueList",
        "value", "levelType",
    ],
    "qualifier": [
        "semanticId", "supplementalSemanticIds", "kind", "type",
        "valueType", "value", "valueId",
    ],
    "resource": ["path", "contentType"],
    "specificAssetId": [
        "semanticId", "supplementalSemanticIds", "name", "value",
        "externalSubjectId",
    ],
    "extension": [
        "semanticId", "supplementalSemanticIds", "name", "valueType",
        "value", "refersTo",
    ],
    "valueReferencePair": ["value", "valueId"],
    "levelType": ["min", "nom", "typ", "max"],
    "langStringTextType": ["language", "text"],
    "langStringNameType": ["language", "text"],
    "langStringPreferredNameTypeIec61360": ["language", "text"],
    "langStringShortNameTypeIec61360": ["language", "text"],
    "langStringDefinitionTypeIec61360": ["language", "text"],
}

DEFAULT_ORDER = list(_SME_HEAD)


# =========================================================
# XML 기본 함수
# =========================================================

def local_name(tag):
    if isinstance(tag, str) and "}" in tag:
        return tag.split("}", 1)[1]
    return tag


def get_namespace(tag):
    if isinstance(tag, str) and tag.startswith("{"):
        return tag.split("}", 1)[0][1:]
    return ""


def ns_tag(element, name):
    namespace = get_namespace(element.tag)
    return f"{{{namespace}}}{name}" if namespace else name


def find_direct_child(element, child_name):
    if element is None:
        return None
    for child in element:
        if local_name(child.tag) == child_name:
            return child
    return None


def find_all_by_name(root, name):
    return [e for e in root.iter() if local_name(e.tag) == name]


def text_of(element):
    if element is None or not element.text:
        return ""
    return element.text.strip()


def get_idshort(element):
    return text_of(find_direct_child(element, "idShort"))


def get_id(element):
    value = text_of(find_direct_child(element, "id"))
    return value or None


def remove_underscore(text):
    return None if text is None else text.replace("_", "")


def insert_ordered(parent, child):
    order = CHILD_ORDER.get(local_name(parent.tag), DEFAULT_ORDER)
    name = local_name(child.tag)

    if name not in order:
        parent.append(child)
        return

    rank = order.index(name)
    index = len(parent)

    for position, existing in enumerate(parent):
        existing_name = local_name(existing.tag)
        existing_rank = (
            order.index(existing_name)
            if existing_name in order
            else len(order) + 1
        )
        if existing_rank > rank:
            index = position
            break

    parent.insert(index, child)


def reorder_children(element):
    order = CHILD_ORDER.get(local_name(element.tag))
    if order is None:
        return False

    children = list(element)
    if len(children) < 2:
        return False

    def rank_of(child):
        name = local_name(child.tag)
        return order.index(name) if name in order else len(order) + 1

    ranks = [rank_of(child) for child in children]
    if all(ranks[i] <= ranks[i + 1] for i in range(len(ranks) - 1)):
        return False

    ordered = [
        child
        for _, _, child in sorted(
            ((rank_of(c), i, c) for i, c in enumerate(children)),
            key=lambda item: (item[0], item[1]),
        )
    ]

    for child in children:
        element.remove(child)
    for child in ordered:
        element.append(child)
    return True


def reorder_all(root, logs):
    count = 0
    for element in list(root.iter()):
        if reorder_children(element):
            count += 1
    if count:
        logs.append({
            "구분": "요소 순서 (XSD sequence)",
            "요소": "(전체)",
            "수정 전": f"{count}개 요소 순서 불일치",
            "수정 후": "스키마 순서로 재정렬",
            "결과": "Auto Fixed",
        })


# =========================================================
# Reference
# =========================================================

def get_reference_value(reference):
    """Reference 의 마지막 key value."""
    keys = find_direct_child(reference, "keys")
    if keys is None:
        return None

    last = None
    for key in keys:
        if local_name(key.tag) != "key":
            continue
        value = text_of(find_direct_child(key, "value"))
        if value:
            last = value
    return last


def get_reference_type(reference):
    return text_of(find_direct_child(reference, "type")) or None


def set_reference(
    reference,
    new_value,
    reference_type="ExternalReference",
    key_type="GlobalReference",
):
    if reference is None:
        return False

    type_element = find_direct_child(reference, "type")
    if type_element is None:
        type_element = ET.Element(ns_tag(reference, "type"))
        insert_ordered(reference, type_element)
    type_element.text = reference_type

    keys = find_direct_child(reference, "keys")
    if keys is None:
        keys = ET.Element(ns_tag(reference, "keys"))
        insert_ordered(reference, keys)

    for child in list(keys):
        keys.remove(child)

    key = ET.SubElement(keys, ns_tag(keys, "key"))
    key_type_element = ET.SubElement(key, ns_tag(key, "type"))
    key_type_element.text = key_type
    value_element = ET.SubElement(key, ns_tag(key, "value"))
    value_element.text = new_value
    return True


def get_semantic_id_value(element):
    return get_reference_value(find_direct_child(element, "semanticId"))


def create_semantic_id(element, reference_value):
    """semanticId 는 그 자체가 reference_t 이다 (<reference> 래퍼 없음)."""
    if find_direct_child(element, "semanticId") is not None:
        return False

    semantic_id = ET.Element(ns_tag(element, "semanticId"))

    type_element = ET.SubElement(semantic_id, ns_tag(semantic_id, "type"))
    type_element.text = "ExternalReference"

    keys = ET.SubElement(semantic_id, ns_tag(semantic_id, "keys"))
    key = ET.SubElement(keys, ns_tag(keys, "key"))
    key_type = ET.SubElement(key, ns_tag(key, "type"))
    key_type.text = "GlobalReference"
    value = ET.SubElement(key, ns_tag(key, "value"))
    value.text = reference_value

    insert_ordered(element, semantic_id)
    return True


def remap_all_keys(root, mapping, logs, label):
    """
    문서 전체의 key/value 를 mapping 에 따라 치환한다.
    Reference 타입(ModelReference/ExternalReference)은 그대로 보존된다.
    """
    if not mapping:
        return 0

    count = 0
    for key in find_all_by_name(root, "key"):
        value_element = find_direct_child(key, "value")
        old = text_of(value_element)
        if old and old in mapping and mapping[old] != old:
            value_element.text = mapping[old]
            count += 1

    if count:
        logs.append({
            "구분": f"{label} 참조 리맵",
            "요소": "(전체)",
            "수정 전": f"{len(mapping)}개 ID 변경",
            "수정 후": f"참조 {count}건 동기화",
            "결과": "Auto Fixed",
        })
    return count


# =========================================================
# idShort 정규화 (KOSMO 규칙)
#
# 대문자로 시작, 영문/숫자/'-'/'_' 사용, '-' 로 끝나지 않음, 2자 이상
# (사용자 요구사항에 따라 '_' 는 제거한다)
# =========================================================

IDSHORT_PATTERN = re.compile(r"^[A-Z][A-Za-z0-9_-]*[A-Za-z0-9_]$")


def sanitize_idshort(value):
    if not value:
        return value

    result = value.replace("_", "")
    allowed = r"[^A-Za-z0-9\-]" if POLICY.get("allow_hyphen", True) \
        else r"[^A-Za-z0-9]"
    result = re.sub(allowed, "", result)
    result = result.rstrip("-")

    if not result:
        return None

    if not result[0].isalpha():
        result = "X" + result

    result = result[0].upper() + result[1:]

    if len(result) < 2:
        result = result + "X"

    return result


def clean_idshorts(root, logs):
    for element in list(root.iter()):
        idshort_element = find_direct_child(element, "idShort")
        if idshort_element is None or not idshort_element.text:
            continue

        old_value = idshort_element.text.strip()
        new_value = sanitize_idshort(old_value)

        if not new_value or old_value == new_value:
            continue

        idshort_element.text = new_value
        reason = (
            "언더바 제거"
            if old_value.replace("_", "") == new_value
            else "KOSMO idShort 규칙"
        )
        logs.append({
            "구분": f"idShort ({reason})",
            "요소": local_name(element.tag),
            "수정 전": old_value,
            "수정 후": new_value,
            "결과": "Auto Fixed",
        })


def clean_display_names(root, logs):
    for display_name in find_all_by_name(root, "displayName"):
        for lang_string in display_name:
            text_element = find_direct_child(lang_string, "text")
            if text_element is None or not text_element.text:
                continue

            old_value = text_element.text
            new_value = remove_underscore(old_value)
            if old_value == new_value:
                continue

            text_element.text = new_value
            logs.append({
                "구분": "displayName",
                "요소": "",
                "수정 전": old_value,
                "수정 후": new_value,
                "결과": "Auto Fixed",
            })


# =========================================================
# Version / Revision
# =========================================================

_NUMERIC = r"\d+(?:\.\d+)*"


def extract_version_revision_from_url(url):
    if not url:
        return None, None

    match = re.search(r"/([^/]+)/([^/]+)$", url.strip().rstrip("/"))
    if not match:
        return None, None

    version, revision = match.group(1), match.group(2)
    if re.fullmatch(_NUMERIC, version) and re.fullmatch(_NUMERIC, revision):
        return version, revision
    return None, None


def get_administration_version_revision(element):
    administration = find_direct_child(element, "administration")
    if administration is None:
        return None, None
    return (
        text_of(find_direct_child(administration, "version")) or None,
        text_of(find_direct_child(administration, "revision")) or None,
    )


def determine_version_revision(element):
    version, revision = get_administration_version_revision(element)
    id_version, id_revision = extract_version_revision_from_url(
        get_id(element)
    )
    version = version or id_version or DEFAULT_VERSION
    revision = revision or id_revision or DEFAULT_REVISION
    return str(version), str(revision)


def ensure_administration(element, version, revision):
    changed = False

    administration = find_direct_child(element, "administration")
    if administration is None:
        administration = ET.Element(ns_tag(element, "administration"))
        insert_ordered(element, administration)
        changed = True

    for name, target in (("version", version), ("revision", revision)):
        child = find_direct_child(administration, name)
        if child is None:
            child = ET.Element(ns_tag(administration, name))
            insert_ordered(administration, child)
            child.text = str(target)
            changed = True
        elif text_of(child) != str(target):
            child.text = str(target)
            changed = True

    return changed


# =========================================================
# URL 생성
# =========================================================

def build_sm_url(equipment_name, idshort, version, revision):
    return (
        f"{BASE_URL}/sm/{remove_underscore(equipment_name)}/"
        f"{remove_underscore(idshort)}/{version}/{revision}"
    )


def get_last_id_token(current_id):
    if not current_id:
        return None
    parts = current_id.rstrip("/").split("/")
    if len(parts) >= 3:
        if (
            re.fullmatch(_NUMERIC, parts[-1])
            and re.fullmatch(_NUMERIC, parts[-2])
        ):
            return parts[-3]
    return parts[-1]


def canonical_cd_name(idshort, old_id=None):
    clean = remove_underscore(idshort)
    token = get_last_id_token(old_id)
    if token in CD_NAME_ALIAS:
        return CD_NAME_ALIAS[token]
    if clean in CD_NAME_ALIAS:
        return CD_NAME_ALIAS[clean]
    return clean


def build_cd_url(idshort, version, revision, old_id=None):
    return (
        f"{BASE_URL}/cd/{canonical_cd_name(idshort, old_id)}/"
        f"{version}/{revision}"
    )


# =========================================================
# AAS / Submodel
# =========================================================

def build_submodel_owner_map(root):
    """submodel id -> AAS idShort. Submodel id 재작성 전에 호출."""
    owner_map = {}
    for shell in find_all_by_name(root, "assetAdministrationShell"):
        shell_name = remove_underscore(get_idshort(shell))
        if not shell_name:
            continue
        container = find_direct_child(shell, "submodels")
        if container is None:
            continue
        for reference in container:
            value = get_reference_value(reference)
            if value:
                owner_map[value] = shell_name
    return owner_map


def get_default_equipment_name(root, xml_path):
    for shell in find_all_by_name(root, "assetAdministrationShell"):
        idshort = get_idshort(shell)
        if idshort:
            return remove_underscore(idshort)

    parts = xml_path.replace("\\", "/").split("/")
    if len(parts) >= 2:
        return remove_underscore(parts[-2])
    return remove_underscore(parts[-1].split(".")[0])


def build_aas_url(idshort, version, revision):
    return (
        f"{BASE_URL}/aas/{remove_underscore(idshort)}/{version}/{revision}"
    )


def fix_aas_administration(root, logs):
    """AAS 의 administration 과 id 를 규격에 맞게 정리하고 id 매핑을 반환."""
    id_map = {}

    for shell in find_all_by_name(root, "assetAdministrationShell"):
        idshort = get_idshort(shell)
        version, revision = determine_version_revision(shell)
        before = get_administration_version_revision(shell)

        if ensure_administration(shell, version, revision):
            logs.append({
                "구분": "AAS Administration",
                "요소": idshort,
                "수정 전": f"{before[0]} / {before[1]}",
                "수정 후": f"{version} / {revision}",
                "결과": "Auto Fixed",
            })

        if not idshort:
            continue

        old_id = get_id(shell)
        new_id = build_aas_url(idshort, version, revision)

        id_element = find_direct_child(shell, "id")
        if id_element is None:
            id_element = ET.Element(ns_tag(shell, "id"))
            insert_ordered(shell, id_element)

        if old_id != new_id:
            if old_id:
                id_map[old_id] = new_id
            id_element.text = new_id
            logs.append({
                "구분": "AAS ID",
                "요소": idshort,
                "수정 전": old_id or "없음",
                "수정 후": new_id,
                "결과": "Auto Fixed",
            })

    return id_map


def fix_submodels(root, owner_map, default_equipment, logs):
    replacement_map = {}
    idshort_map = {}

    for submodel in find_all_by_name(root, "submodel"):
        idshort = get_idshort(submodel)
        if not idshort:
            logs.append({
                "구분": "Submodel",
                "요소": "(idShort 없음)",
                "수정 전": get_id(submodel) or "",
                "수정 후": "건너뜀",
                "결과": "Warning",
            })
            continue

        old_id = get_id(submodel) or ""
        equipment = owner_map.get(old_id) or default_equipment
        version, revision = determine_version_revision(submodel)
        new_url = build_sm_url(equipment, idshort, version, revision)

        id_element = find_direct_child(submodel, "id")
        if id_element is None:
            id_element = ET.Element(ns_tag(submodel, "id"))
            insert_ordered(submodel, id_element)

        if old_id != new_url:
            if old_id:
                replacement_map[old_id] = new_url
            id_element.text = new_url
            logs.append({
                "구분": "Submodel ID",
                "요소": idshort,
                "수정 전": old_id or "없음",
                "수정 후": new_url,
                "결과": "Auto Fixed",
            })
        idshort_map[idshort] = new_url

        semantic_id = find_direct_child(submodel, "semanticId")
        if semantic_id is not None and is_standard_id(
            get_reference_value(semantic_id)
        ):
            # IDTA 템플릿 Submodel: semanticId 원본 유지
            ensure_administration(submodel, version, revision)
            fix_submodel_kind(submodel, idshort, logs)
            continue
        if semantic_id is None:
            create_semantic_id(submodel, new_url)
            logs.append({
                "구분": "Submodel SemanticID",
                "요소": idshort,
                "수정 전": "없음",
                "수정 후": new_url,
                "결과": "Auto Fixed",
            })
        else:
            old_semantic = get_reference_value(semantic_id)
            if old_semantic != new_url:
                set_reference(semantic_id, new_url)
                logs.append({
                    "구분": "Submodel SemanticID",
                    "요소": idshort,
                    "수정 전": old_semantic or "없음",
                    "수정 후": new_url,
                    "결과": "Auto Fixed",
                })

        ensure_administration(submodel, version, revision)
        fix_submodel_kind(submodel, idshort, logs)

    return replacement_map, idshort_map


SUBMODEL_KIND = {"value": "Template"}

# 정책 (UI 옵션으로 덮어쓴다)
POLICY = {
    # 표준 id CD(IDTA/IRDI)의 dataType 등 내용 보호
    "protect_standard_cd": True,
    # 표준 IRI semanticId 에 CD 가 없을 때: "smartfactory" | "same"
    "template_iri_strategy": "smartfactory",
    # SML 항목 idShort 제거: "auto"(3.0 만) | "always" | "never"
    "sml_item_idshort": "auto",
    # assetKind 강제값 (None 이면 건드리지 않음)
    "asset_kind": "Type",
    # globalAssetId 를 smart-factory 형식으로 맞춤
    "fix_global_asset_id": True,
    # Operation/BasicEventElement/Capability 제거 (False 면 경고만)
    "remove_special_sme": False,
    # 한글 사전 값으로 기존 텍스트 덮어쓰기
    "glossary_overwrite": False,
    # IDTA/표준 CD 에 한글(이름·정의)만 추가 (영문 원본은 유지)
    "standard_cd_korean": True,
    # 내부용: 네임스페이스에 따라 run_rules 가 설정
    "allow_hyphen": True,
}

# 한글/영문 사전: idShort -> {"ko_name","ko_definition","en_name","en_definition"}
GLOSSARY = {}
# CD id -> {...} (검토 엑셀에서 읽은 값, idShort 사전보다 우선)
GLOSSARY_BY_ID = {}


def is_protected_element(element):
    """semanticId 가 표준(IDTA/IRDI) 이면 템플릿 원본 요소로 보고 보호."""
    return is_standard_id(get_semantic_id_value(element))


def is_protected_cd(cd):
    return POLICY.get("protect_standard_cd", True) and is_standard_id(
        get_id(cd)
    )


def fix_submodel_kind(submodel, idshort, logs):
    """KOSMO-Submodel #4: 참조모델의 Submodel kind 는 Template."""
    target = SUBMODEL_KIND.get("value")
    if not target:
        return
    kind = find_direct_child(submodel, "kind")
    old = text_of(kind)
    if old == target:
        return
    if kind is None:
        kind = ET.Element(ns_tag(submodel, "kind"))
        insert_ordered(submodel, kind)
    kind.text = target
    logs.append({
        "구분": "Submodel kind (KOSMO-Submodel #4)",
        "요소": idshort,
        "수정 전": old or "없음",
        "수정 후": target,
        "결과": "Auto Fixed",
    })


def normalize_aas_submodel_references(root, idshort_map, logs):
    """AAS 의 submodel 참조를 ModelReference/Submodel 형태로 맞춘다."""
    valid_ids = set(idshort_map.values())

    for shell in find_all_by_name(root, "assetAdministrationShell"):
        shell_name = get_idshort(shell)
        container = find_direct_child(shell, "submodels")
        if container is None:
            continue

        for reference in container:
            value = get_reference_value(reference)
            if not value:
                continue

            new_value = value
            if value not in valid_ids:
                token = remove_underscore(get_last_id_token(value) or "")
                if token in idshort_map:
                    new_value = idshort_map[token]
                else:
                    logs.append({
                        "구분": "AAS Submodel Reference",
                        "요소": shell_name,
                        "수정 전": value,
                        "수정 후": "대응 Submodel 없음 - 수동 확인 필요",
                        "결과": "Warning",
                    })
                    continue

            keys = find_direct_child(reference, "keys")
            key_count = (
                len([k for k in keys if local_name(k.tag) == "key"])
                if keys is not None else 0
            )
            current_type = get_reference_type(reference)

            if (
                value == new_value
                and current_type == "ModelReference"
                and key_count == 1
            ):
                continue

            set_reference(
                reference, new_value,
                reference_type="ModelReference", key_type="Submodel",
            )
            logs.append({
                "구분": "AAS Submodel Reference",
                "요소": shell_name,
                "수정 전": f"{current_type} / {value}",
                "수정 후": f"ModelReference / {new_value}",
                "결과": "Auto Fixed",
            })


# =========================================================
# ConceptDescription
# =========================================================

def get_concept_descriptions(root):
    return find_all_by_name(root, "conceptDescription")


def fix_cd_missing_idshort(root, logs):
    """KOSMO-CD #1: idShort 가 없는 CD 에 preferredName(en) 기반 idShort 부여."""
    for cd in get_concept_descriptions(root):
        if get_idshort(cd):
            continue
        content = get_iec61360_content(cd)
        name = None
        if content is not None:
            for lang_string in find_direct_child(content, "preferredName") \
                    or []:
                if text_of(find_direct_child(lang_string, "language")) \
                        .lower().startswith("en"):
                    words = re.findall(
                        r"[A-Za-z0-9]+",
                        text_of(find_direct_child(lang_string, "text")),
                    )
                    name = "".join(w[:1].upper() + w[1:] for w in words)
                    break
        name = sanitize_idshort(name) if name else name_from_semantic_id(
            get_id(cd)
        )
        if not name:
            logs.append({
                "구분": "CD idShort 없음 (KOSMO-CD #1)",
                "요소": get_id(cd),
                "수정 전": "idShort / preferredName(en) 없음",
                "수정 후": "이름 추론 불가 - 수동 입력 필요",
                "결과": "Warning",
            })
            continue
        element = ET.Element(ns_tag(cd, "idShort"))
        element.text = name
        insert_ordered(cd, element)
        logs.append({
            "구분": "CD idShort 보충 (KOSMO-CD #1)",
            "요소": get_id(cd),
            "수정 전": "없음",
            "수정 후": name,
            "결과": "Auto Fixed",
        })


def normalize_concept_descriptions(root, logs):
    """
    CD id 를 /ids/cd/{name}/{version}/{revision} 로 통일하고
    old_id -> new_id 매핑을 반환한다.
    """
    id_map = {}

    for cd in get_concept_descriptions(root):
        idshort = get_idshort(cd)
        if not idshort:
            continue

        old_id = get_id(cd)
        if is_standard_id(old_id):
            # IDTA 템플릿 / 표준사전 CD 는 id 를 바꾸지 않는다
            continue
        version, revision = determine_version_revision(cd)
        new_id = build_cd_url(idshort, version, revision, old_id)

        id_element = find_direct_child(cd, "id")
        if id_element is None:
            id_element = ET.Element(ns_tag(cd, "id"))
            insert_ordered(cd, id_element)

        if old_id != new_id:
            if old_id:
                id_map[old_id] = new_id
            id_element.text = new_id
            logs.append({
                "구분": "ConceptDescription ID",
                "요소": idshort,
                "수정 전": old_id or "없음",
                "수정 후": new_id,
                "결과": "Auto Fixed",
            })

        ensure_administration(cd, version, revision)

    return id_map


def remove_duplicate_concept_descriptions(root, logs):
    for parent in list(root.iter()):
        seen = set()
        for child in list(parent):
            if local_name(child.tag) != "conceptDescription":
                continue
            key = (get_idshort(child), get_id(child))
            if key not in seen:
                seen.add(key)
                continue
            parent.remove(child)
            logs.append({
                "구분": "Duplicate CD",
                "요소": key[0],
                "수정 전": key[1] or "",
                "수정 후": "중복 1건 제거",
                "결과": "Auto Fixed",
            })


def get_concept_descriptions_container(root):
    container = find_direct_child(root, "conceptDescriptions")
    if container is None:
        container = ET.Element(ns_tag(root, "conceptDescriptions"))
        insert_ordered(root, container)
    return container


def build_cd_index(root):
    """(id 집합, idShort -> id (유일한 경우만), idShort 집합)"""
    ids = set()
    by_idshort = {}
    all_idshorts = set()

    for cd in get_concept_descriptions(root):
        cd_id = get_id(cd)
        idshort = get_idshort(cd)
        if cd_id:
            ids.add(cd_id)
        if idshort:
            all_idshorts.add(idshort)
        if cd_id and idshort:
            by_idshort.setdefault(idshort, set()).add(cd_id)

    unique = {
        name: next(iter(values))
        for name, values in by_idshort.items()
        if len(values) == 1
    }
    return ids, unique, all_idshorts


# =========================================================
# SME <-> CD semanticId 해소
# =========================================================

SME_TYPES = {
    "property", "multiLanguageProperty", "range", "file", "blob",
    "referenceElement", "relationshipElement", "annotatedRelationshipElement",
    "entity", "operation", "capability", "basicEventElement",
    "submodelElementCollection", "submodelElementList",
}

SME_DEFAULT_DATATYPE = {
    "multiLanguageProperty": "STRING_TRANSLATABLE",
    "file": "FILE",
    "blob": "BLOB",
    "referenceElement": "IRI",
}


def name_from_semantic_id(value):
    """semanticId URL/IRDI 에서 CD 이름 후보를 만든다."""
    if not value:
        return None

    if is_irdi(value):
        return None  # IRDI 에서는 의미 있는 이름을 만들 수 없다
    token = get_last_id_token(value)
    candidate = sanitize_idshort(token)
    if (
        candidate
        and len(candidate) >= 2
        and re.search(r"[A-Za-z]{2}", candidate)
    ):
        return candidate
    return None


def unique_cd_name(base, used):
    name = base or "GeneratedConcept"
    if name not in used:
        used.add(name)
        return name
    index = 2
    while f"{name}{index}" in used:
        index += 1
    result = f"{name}{index}"
    used.add(result)
    return result


def create_concept_description(
    root, name, cd_id, datatype, used_names, administration=True
):
    container = get_concept_descriptions_container(root)
    cd = ET.SubElement(container, ns_tag(container, "conceptDescription"))

    idshort = ET.SubElement(cd, ns_tag(cd, "idShort"))
    idshort.text = name

    display_name = ET.SubElement(cd, ns_tag(cd, "displayName"))
    add_lang_string(display_name, "langStringNameType", name)
    add_lang_string(
        display_name, "langStringNameType", preferred_name_ko_for(name), "ko"
    )
    description = ET.SubElement(cd, ns_tag(cd, "description"))
    add_lang_string(
        description, "langStringTextType", description_text_for(name)
    )
    add_lang_string(
        description, "langStringTextType", description_text_ko_for(name), "ko"
    )

    if administration:
        admin = ET.SubElement(cd, ns_tag(cd, "administration"))
        version = ET.SubElement(admin, ns_tag(admin, "version"))
        version.text = DEFAULT_VERSION
        revision = ET.SubElement(admin, ns_tag(admin, "revision"))
        revision.text = DEFAULT_REVISION

    id_element = ET.SubElement(cd, ns_tag(cd, "id"))
    id_element.text = cd_id

    build_iec61360(cd, name, datatype)
    used_names.add(name)
    return cd


def make_reference_element(parent, tag, value):
    """parent 아래 reference_t 형태(tag)를 만들어 반환 (insert 는 호출측)."""
    ref = ET.Element(ns_tag(parent, tag))
    set_reference(ref, value)
    return ref


def add_supplemental_semantic_id(element, value):
    container = find_direct_child(element, "supplementalSemanticIds")
    if container is None:
        container = ET.Element(ns_tag(element, "supplementalSemanticIds"))
        insert_ordered(element, container)
    for ref in container:
        if get_reference_value(ref) == value:
            return False
    ref = ET.Element(ns_tag(container, "reference"))
    set_reference(ref, value)
    container.insert(0, ref)
    return True


def add_is_case_of(cd, value):
    container = find_direct_child(cd, "isCaseOf")
    if container is None:
        container = ET.Element(ns_tag(cd, "isCaseOf"))
        insert_ordered(cd, container)
    for ref in container:
        if get_reference_value(ref) == value:
            return
    ref = ET.SubElement(container, ns_tag(container, "reference"))
    set_reference(ref, value)


def find_cd_by_idshort(root, name):
    for cd in get_concept_descriptions(root):
        if get_idshort(cd) == name:
            return cd
    return None


def clone_cd_with_id(root, source, new_id, name):
    """source CD 내용을 복사해 새 id 의 CD 를 만든다."""
    container = get_concept_descriptions_container(root)
    clone = copy.deepcopy(source)
    id_element = find_direct_child(clone, "id")
    id_element.text = new_id
    idshort = find_direct_child(clone, "idShort")
    if idshort is not None:
        idshort.text = name
    container.append(clone)
    return clone


def resolve_semantic_ids(root, logs, create_missing=True):
    """
    1) semanticId 가 이미 CD id 와 매칭되면 그대로 둔다.
    2) 표준 semanticId(IDTA/IRDI) 인데 CD 가 없으면
       - IRDI: 같은 id 로 CD 를 만든다 (semanticId 는 그대로).
       - IRI : smart-factory CD 를 만들어 연결하고, 원래 IRI 는
               supplementalSemanticId + CD isCaseOf 로 남긴다.
       (같은 idShort 의 CD 가 있으면 그 내용을 복사한다)
    3) 자체 semanticId 가 매칭되지 않으면 같은 idShort 의 CD 로 연결.
    4) 그래도 없으면 CD 를 새로 만들어 연결.
    """
    cd_ids, cd_by_idshort, used_names = build_cd_index(root)

    orphan_groups = {}
    standard_groups = {}
    no_semantic = []

    for element in list(root.iter()):
        if local_name(element.tag) not in SME_TYPES:
            continue

        idshort = get_idshort(element)
        semantic_id = find_direct_child(element, "semanticId")

        if semantic_id is None:
            no_semantic.append(element)
            continue

        value = get_reference_value(semantic_id)

        if value and value in cd_ids:
            continue

        if is_standard_id(value):
            standard_groups.setdefault(value, []).append((element, idshort))
            continue

        if idshort and idshort in cd_by_idshort:
            target = cd_by_idshort[idshort]
            set_reference(semantic_id, target)
            logs.append({
                "구분": "SME ↔ CD SemanticID",
                "요소": idshort,
                "수정 전": value or "없음",
                "수정 후": target,
                "결과": "Auto Fixed",
            })
            continue

        orphan_groups.setdefault(value or "", []).append((element, idshort))

    # --- 표준 semanticId (템플릿 요소) ---
    for value, members in standard_groups.items():
        idshorts = [n for _, n in members if n]
        name = (
            (sanitize_idshort(idshorts[0]) if idshorts else None)
            or name_from_semantic_id(value)
            or "GeneratedConcept"
        )
        if not create_missing:
            logs.append({
                "구분": "SemanticID 미매칭 (템플릿)",
                "요소": name,
                "수정 전": value,
                "수정 후": f"대응 CD 없음 ({len(members)}개 요소)",
                "결과": "Warning",
            })
            continue

        source = find_cd_by_idshort(root, name)
        strategy = POLICY.get("template_iri_strategy", "smartfactory")

        if is_irdi(value) or strategy == "same":
            if source is not None:
                clone_cd_with_id(root, source, value, name)
                how = f"CD '{name}' 내용 복사"
            else:
                create_concept_description(
                    root, name, value, None, used_names, administration=False
                )
                how = "골격 생성"
            cd_ids.add(value)
            logs.append({
                "구분": "CD 생성 (템플릿 semanticId 유지)",
                "요소": name,
                "수정 전": f"{value} 에 대응 CD 없음",
                "수정 후": f"같은 id 로 CD 생성 ({how})",
                "결과": "Auto Fixed",
            })
            continue

        # IRI: smart-factory CD + supplementalSemanticId + isCaseOf
        new_id = build_cd_url(name, DEFAULT_VERSION, DEFAULT_REVISION)
        if new_id not in cd_ids:
            if source is not None:
                cd = clone_cd_with_id(root, source, new_id, name)
                ensure_administration(cd, DEFAULT_VERSION, DEFAULT_REVISION)
            else:
                cd = create_concept_description(
                    root, name, new_id, None, used_names
                )
            add_is_case_of(cd, value)
            cd_ids.add(new_id)
        for element, _ in members:
            set_reference(find_direct_child(element, "semanticId"), new_id)
            add_supplemental_semantic_id(element, value)
        logs.append({
            "구분": "CD 생성 (템플릿 IRI → smart-factory)",
            "요소": name,
            "수정 전": value,
            "수정 후": f"{new_id} (원래 IRI는 supplementalSemanticId·isCaseOf 로 유지)",
            "결과": "Auto Fixed",
        })

    # --- semanticId 자체가 없는 요소 ---
    for element in no_semantic:
        idshort = get_idshort(element)
        if idshort and idshort in cd_by_idshort:
            target = cd_by_idshort[idshort]
            create_semantic_id(element, target)
            logs.append({
                "구분": "Missing SemanticID",
                "요소": idshort,
                "수정 전": "없음",
                "수정 후": target,
                "결과": "Auto Fixed",
            })
        else:
            orphan_groups.setdefault(f"__NOSEM__{idshort}", []).append(
                (element, idshort)
            )

    if not create_missing:
        for value, members in orphan_groups.items():
            logs.append({
                "구분": "SemanticID 미매칭",
                "요소": members[0][1] or "(idShort 없음)",
                "수정 전": value,
                "수정 후": f"대응 CD 없음 ({len(members)}개 요소)",
                "결과": "Warning",
            })
        return

    for value, members in orphan_groups.items():
        idshorts = [name for _, name in members if name]
        base_name = None

        if idshorts:
            base_name = sanitize_idshort(idshorts[0])
        if not base_name and not value.startswith("__NOSEM__"):
            base_name = name_from_semantic_id(value)
        if not base_name:
            base_name = "GeneratedConcept"

        name = unique_cd_name(base_name, used_names)

        element_type = local_name(members[0][0].tag)
        datatype = SME_DEFAULT_DATATYPE.get(element_type)
        if element_type in ("property", "range"):
            value_type = text_of(
                find_direct_child(members[0][0], "valueType")
            )
            datatype = preferred_cd_datatype(value_type, False) or "STRING"

        new_id = build_cd_url(name, DEFAULT_VERSION, DEFAULT_REVISION)
        create_concept_description(root, name, new_id, datatype, used_names)
        cd_ids.add(new_id)
        cd_by_idshort.setdefault(name, new_id)

        for element, _ in members:
            semantic_id = find_direct_child(element, "semanticId")
            if semantic_id is None:
                create_semantic_id(element, new_id)
            else:
                set_reference(semantic_id, new_id)

        logs.append({
            "구분": "CD 신규 생성 - 정의 작성 필요",
            "요소": name,
            "수정 전": (
                "semanticId 없음"
                if value.startswith("__NOSEM__")
                else value
            ),
            "수정 후": f"{new_id} ({len(members)}개 요소 연결)",
            "결과": "Auto Fixed",
        })


# =========================================================
# Property valueType / value
# =========================================================

def split_tokens(idshort):
    text = re.sub(r"[_\-\s]+", " ", idshort or "")
    text = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", text)
    text = re.sub(r"(?<=[A-Z])(?=[A-Z][a-z])", " ", text)
    return [t.lower() for t in text.split() if t]


DATETIME_TOKENS = {"datetime", "timestamp"}
DATETIME_CONTEXT = {"start", "end", "failure", "event", "issue", "expiry"}

# ---------------------------------------------------------
# 타입 판정은 "head word"(복합명사의 마지막 단어)를 우선으로 본다.
# 영어 복합명사는 맨 뒤 단어가 그 항목의 정체를 결정한다.
#   TensionControlMethod -> Method  -> 문자열
#   TargetTemperature    -> Temperature -> 실수
#   MaximumTurnCount     -> Count   -> 정수
# "of" 가 있으면 그 앞부분이 본체다. (YearOfConstruction -> Year)
# ---------------------------------------------------------

# 문자열이어야 하는 head
STRING_HEADS = {
    # 식별자·코드
    "id", "ids", "identifier", "number", "no", "code", "serial",
    "version", "revision", "barcode", "batch", "lot", "sku", "gtin",
    "mac", "ip", "key", "password", "checksum", "hash", "signature",
    "uuid", "guid", "reference", "path", "filename", "extension",
    # 이름·설명
    "name", "text", "description", "comment", "note", "label", "title",
    "designation", "remark", "reason", "message", "caption", "summary",
    # 분류·상태
    "type", "kind", "class", "category", "grade", "shape", "material",
    "color", "colour", "state", "status", "mode", "method", "action",
    "priority", "role", "function", "protocol", "encoding",
    "standard", "specification", "system", "format", "unit", "symbol",
    "component", "recipe", "domain", "organization", "department",
    "source", "subject", "scope", "phase", "stage",
    # 주소·연락처
    "address", "street", "city", "country", "region", "zip", "postal",
    "phone", "fax", "email", "contact", "language", "currency",
    "location", "site", "plant", "line", "area",
    # 기타
    "year", "declaration", "certificate", "approval", "guideline",
}

# 참/거짓 head
BOOLEAN_HEADS = {
    "flag", "enabled", "disabled", "detected", "available", "active",
    "valid", "installed", "connected", "required", "present", "yesno",
    "primary", "used", "ok", "pass", "fail",
}

# 정수 head (개수·횟수)
INTEGER_HEADS = {
    "count", "quantity", "qty", "retry", "retries", "cycles", "attempts",
    "errors", "warnings", "alarms", "channels", "slots", "pieces",
    "units", "shots", "passes", "iterations", "turns", "layers",
    "occurrences", "hits", "faults",
}

# 실수 head (측정값·물리량·지표)
REAL_HEADS = {
    # 물리량
    "temperature", "temp", "pressure", "speed", "velocity", "current",
    "voltage", "power", "frequency", "flow", "flowrate", "length",
    "width", "height", "diameter", "radius", "distance", "thickness",
    "depth", "weight", "mass", "volume", "torque", "load", "force",
    "energy", "consumption", "angle", "amplitude", "humidity",
    "density", "viscosity", "resistance", "impedance", "inductance",
    "capacitance", "flux", "luminance", "noise", "wavelength",
    "intensity", "gain", "altitude", "acceleration", "vibration",
    "tension", "pitch", "feed", "stroke", "clearance", "gap",
    "backlash", "runout", "roughness", "hardness", "conductivity",
    "capacity", "throughput", "bandwidth", "latency", "rpm",
    # 비율·정확도·지표
    "accuracy", "precision", "tolerance", "deviation", "error",
    "ratio", "rate", "efficiency", "score", "factor", "index",
    "probability", "percentage", "percent", "offset", "variance",
    "resolution", "calibration", "coordinate", "position", "setpoint",
    "threshold", "limit", "margin", "span", "wear", "degradation",
    "lifetime", "rul", "mtbf", "mttr", "utilization",
    "latitude", "longitude", "elevation", "delta", "difference",
    "average", "mean", "median", "sum", "total",
}

# head 가 애매할 때 숫자임을 시사하는 접두어
NUMERIC_PREFIXES = {
    "minimum", "maximum", "min", "max", "average", "avg", "mean",
    "total", "sum", "rated", "nominal", "peak", "median", "actual",
    "measured", "calculated", "estimated", "remaining", "elapsed",
}

BOOLEAN_PREFIXES = ("is", "has", "can", "was", "does", "are", "should")

URI_HEADS = {"uri", "url", "link", "href", "endpoint", "website"}

STATUS_TOKENS = {"status", "state", "mode"}
RESULT_TOKENS = {"result", "judgement", "judgment", "verdict"}

DECIMAL_PATTERN = re.compile(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?")
INTEGER_PATTERN = re.compile(r"[+-]?\d+")


def looks_like_padded_number(text):
    """0044 처럼 앞자리 0이 있는 값은 문자열로 봐야 한다."""
    stripped = text.lstrip("+-")
    return len(stripped) > 1 and stripped.startswith("0") and "." not in stripped


# 그 자체로는 정체를 알려주지 않는 수식어성 head.
# 이런 단어가 끝에 오면 한 칸 앞 단어를 본다.
#   TurnCountTarget -> Count -> 정수
#   AssessmentTarget -> Assessment -> (미확정)
MODIFIER_HEADS = {
    "target", "value", "setting", "spec", "actual", "measured",
    "new", "old", "previous", "next",
}


def head_tokens(idshort):
    """
    idShort 를 토큰으로 쪼개고, 'of' 앞부분을 본체로 삼는다.
    수식어성 단어가 끝에 오면 한 칸 앞으로 넘어간다.
    반환: (전체 토큰 리스트, head 단어)
    """
    tokens = split_tokens(idshort)
    if not tokens:
        return [], ""

    core = tokens
    for separator in ("of", "for", "in", "at", "per"):
        if separator in tokens:
            before = tokens[: tokens.index(separator)]
            if before:
                core = before
                break

    index = len(core) - 1
    while index > 0 and core[index] in MODIFIER_HEADS:
        index -= 1

    return tokens, core[index]


def infer_value_type(idshort, value_text=None):
    """
    idShort 의 head word 와 실제 값에서 xsd valueType 을 추론한다.

    판정 순서
      1. Is/Has 로 시작하면 참/거짓
      2. head word (복합명사의 마지막 단어)
      3. 실제 값의 형태
      4. 전체 토큰 / 숫자 접두어
    """
    tokens, head = head_tokens(idshort)
    token_set = set(tokens)
    text = (value_text or "").strip()

    # 1. Is/Has/Can 접두, 또는 값이 명확히 true/false
    if tokens and tokens[0] in BOOLEAN_PREFIXES and len(tokens) > 1:
        return "xs:boolean"
    if text.lower() in {"true", "false"}:
        return "xs:boolean"

    # 2. head word 기반
    if head in STRING_HEADS:
        return "xs:string"
    if head in BOOLEAN_HEADS:
        return "xs:boolean"
    if head in URI_HEADS:
        return "xs:anyURI"

    if head in DATETIME_TOKENS:
        return "xs:dateTime"
    if head == "date":
        return "xs:dateTime" if token_set & {"datetime"} else "xs:date"
    if head == "time":
        # StartTime/EndTime 은 시각, ResponseTime/CycleTime 은 소요 시간
        if token_set & DATETIME_CONTEXT:
            return "xs:dateTime"
        return "xs:float"

    if head in INTEGER_HEADS:
        return "xs:integer"
    if head in REAL_HEADS:
        return "xs:float"

    # 3. 실제 값의 형태
    if text:
        if looks_like_padded_number(text):
            return "xs:string"
        if text.lower() in {"true", "false"}:
            return "xs:boolean"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return "xs:date"
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}T[\d:.]+.*", text):
            return "xs:dateTime"
        if INTEGER_PATTERN.fullmatch(text):
            return "xs:float" if tokens and tokens[0] in NUMERIC_PREFIXES \
                else "xs:integer"
        if DECIMAL_PATTERN.fullmatch(text):
            return "xs:float"
        return "xs:string"

    # 4. 전체 토큰 / 숫자 접두어
    if token_set & BOOLEAN_HEADS:
        return "xs:boolean"
    if token_set & DATETIME_TOKENS:
        return "xs:dateTime"
    if "date" in token_set:
        return "xs:date"
    if token_set & INTEGER_HEADS:
        return "xs:integer"
    if token_set & REAL_HEADS:
        return "xs:float"
    if tokens and tokens[0] in NUMERIC_PREFIXES:
        return "xs:float"

    return "xs:string"


def infer_property_type_and_value(idshort):
    """valueType 과 채워 넣을 샘플 값을 함께 돌려준다."""
    value_type = infer_value_type(idshort)

    if value_type == "xs:string":
        token_set = set(split_tokens(idshort))
        if token_set & STATUS_TOKENS:
            return "xs:string", "Normal"
        if token_set & RESULT_TOKENS:
            return "xs:string", "Success"
        return "xs:string", "Sample"

    return value_type, sample_value_for(value_type)


def normalize_value_type(value_type):
    if not value_type:
        return ""
    return value_type.replace(
        "http://www.w3.org/2001/XMLSchema#", "xs:"
    ).strip()


STRING_TYPES = {"xs:string", "string", "xs:normalizedString"}
URI_TYPES = {"xs:anyURI", "anyURI"}
REAL_TYPES = {
    "xs:float", "xs:double", "xs:decimal", "float", "double", "decimal",
}
INTEGER_TYPES = {
    "xs:int", "xs:integer", "xs:long", "xs:short", "xs:byte",
    "xs:unsignedInt", "xs:unsignedLong", "xs:unsignedShort",
    "xs:unsignedByte", "xs:nonNegativeInteger", "xs:positiveInteger",
    "xs:negativeInteger", "xs:nonPositiveInteger",
    "int", "integer", "long", "short",
}
BOOLEAN_TYPES = {"xs:boolean", "boolean"}


SAMPLE_BY_TYPE = {
    "xs:dateTime": SAMPLE_DATETIME,
    "xs:date": SAMPLE_DATE,
    "xs:time": SAMPLE_TIME,
    "xs:boolean": "true",
    "xs:anyURI": "https://www.smart-factory.kr/",
}


def sample_value_for(value_type, fallback="Sample"):
    normalized = normalize_value_type(value_type)
    if normalized in SAMPLE_BY_TYPE:
        return SAMPLE_BY_TYPE[normalized]
    if normalized in REAL_TYPES:
        return "1.0"
    if normalized in INTEGER_TYPES:
        return "1"
    return fallback


def value_matches_type(value_text, value_type):
    """빈 값은 항상 호환으로 본다(별도 단계에서 채운다)."""
    if value_text is None or not value_text.strip():
        return True

    normalized = normalize_value_type(value_type)
    text = value_text.strip()

    if normalized in BOOLEAN_TYPES:
        return text.lower() in {"true", "false", "1", "0"}
    if normalized in INTEGER_TYPES:
        return re.fullmatch(r"[+-]?\d+", text) is not None
    if normalized in REAL_TYPES:
        return re.fullmatch(
            r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?", text
        ) is not None
    if normalized == "xs:date":
        return re.fullmatch(
            r"\d{4}-\d{2}-\d{2}(Z|[+-]\d{2}:\d{2})?", text
        ) is not None
    if normalized == "xs:dateTime":
        return re.fullmatch(
            r"\d{4}-\d{2}-\d{2}T[\d:.]+(Z|[+-]\d{2}:\d{2})?", text
        ) is not None
    if normalized == "xs:time":
        return re.fullmatch(
            r"\d{2}:\d{2}(:\d{2}(\.\d+)?)?(Z|[+-]\d{2}:\d{2})?", text
        ) is not None
    return True


def infer_type_from_value(value_text):
    """실제 값의 형태에서 xsd 타입을 추론한다."""
    if value_text is None:
        return None
    text = value_text.strip()
    if not text:
        return None

    if text.lower() in {"true", "false"}:
        return "xs:boolean"
    if re.fullmatch(r"[+-]?\d+", text):
        return "xs:integer"
    if re.fullmatch(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?", text):
        return "xs:float"
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(Z|[+-]\d{2}:\d{2})?", text):
        return "xs:date"
    if re.fullmatch(
        r"\d{4}-\d{2}-\d{2}T[\d:.]+(Z|[+-]\d{2}:\d{2})?", text
    ):
        return "xs:dateTime"
    return "xs:string"


def normalize_value_text(value_text, value_type):
    """xs:boolean 은 소문자여야 한다."""
    if value_text is None:
        return None
    normalized = normalize_value_type(value_type)
    if normalized in BOOLEAN_TYPES:
        lowered = value_text.strip().lower()
        if lowered in {"true", "1"}:
            return "true"
        if lowered in {"false", "0"}:
            return "false"
    return value_text


DOUBLE_ALIASES = {
    "xs:double", "xs:decimal", "double", "decimal",
    "http://www.w3.org/2001/XMLSchema#double",
    "http://www.w3.org/2001/XMLSchema#decimal",
}


def unify_real_value_types(root, logs):
    """실수 타입을 xs:float 하나로 통일한다 (xs:double / xs:decimal 포함)."""
    count = 0
    for element_name in ("property", "range"):
        for element in find_all_by_name(root, element_name):
            if is_protected_element(element):
                continue
            value_type_element = find_direct_child(element, "valueType")
            if value_type_element is None:
                continue
            current = text_of(value_type_element)
            if normalize_value_type(current) not in DOUBLE_ALIASES:
                continue
            value_type_element.text = "xs:float"
            count += 1

    # SubmodelElementList 의 valueTypeListElement 도 함께 맞춘다
    for sml in find_all_by_name(root, "submodelElementList"):
        vt_element = find_direct_child(sml, "valueTypeListElement")
        if vt_element is None:
            continue
        if normalize_value_type(text_of(vt_element)) in DOUBLE_ALIASES:
            vt_element.text = "xs:float"
            count += 1

    if count:
        logs.append({
            "구분": "실수 타입 통일",
            "요소": "(전체)",
            "수정 전": "xs:double / xs:decimal",
            "수정 후": f"xs:float ({count}건)",
            "결과": "Auto Fixed",
        })


def refine_property_value_types(root, logs):
    """
    valueType 이 범용 xs:string 으로만 되어 있는 Property 를
    idShort 의 의미와 실제 값을 근거로 구체적인 타입으로 바꾼다.

    - 이미 구체적인 타입(xs:int, xs:float ...)은 건드리지 않는다.
    - SerialNumber, OrderNumber, 0044 처럼 문자열이어야 하는 값은 그대로 둔다.
    - 값을 새로 만들지는 않는다. 기존 값이 있으면 그 형태를 근거로 삼는다.
    """
    for prop in find_all_by_name(root, "property"):
        value_type_element = find_direct_child(prop, "valueType")
        current = normalize_value_type(text_of(value_type_element))

        # 값이 없거나 이미 구체적인 타입이면 여기서는 건드리지 않는다
        if current and current not in STRING_TYPES:
            continue

        idshort = get_idshort(prop)
        if not idshort or is_protected_element(prop):
            continue

        value_text = text_of(find_direct_child(prop, "value"))
        inferred = infer_value_type(idshort, value_text)

        if inferred == current or normalize_value_type(inferred) in STRING_TYPES:
            continue

        # 기존 값이 새 타입과 맞지 않으면 바꾸지 않는다 (값 보존 우선)
        if value_text and not value_matches_type(value_text, inferred):
            continue

        if value_type_element is None:
            value_type_element = ET.Element(ns_tag(prop, "valueType"))
            insert_ordered(prop, value_type_element)

        value_type_element.text = inferred
        logs.append({
            "구분": "Property valueType 세분화",
            "요소": idshort,
            "수정 전": f"{current or '없음'} / value={value_text[:20] or '없음'}",
            "수정 후": inferred,
            "결과": "Auto Fixed",
        })

        # 값 표기 정규화 (True -> true)
        value_element = find_direct_child(prop, "value")
        if value_element is not None and value_text:
            normalized = normalize_value_text(value_text, inferred)
            if normalized != value_text:
                value_element.text = normalized
                logs.append({
                    "구분": "Property value 표기 정규화",
                    "요소": idshort,
                    "수정 전": value_text,
                    "수정 후": normalized,
                    "결과": "Auto Fixed",
                })


# 이번 실행에서 임의로 채운 예시값 (요소 경로)
AUTO_VALUES = set()

# 재실행 시에도 임의값으로 보이는 값
SAMPLE_LIKE = {
    "Sample", "Normal", "Success", SAMPLE_DATE, SAMPLE_DATETIME,
    SAMPLE_TIME, "https://www.smart-factory.kr/",
}


def element_path(element, parent_map):
    """Submodel idShort 부터의 경로. SML 항목은 [순번] 으로 표시."""
    parts = []
    current = element
    while current is not None:
        name = local_name(current.tag)
        if name == "submodel":
            parts.append(get_idshort(current) or "(submodel)")
            break
        if name in SME_TYPES:
            idshort = get_idshort(current)
            if idshort:
                parts.append(idshort)
            else:
                container = parent_map.get(current)
                siblings = [
                    c for c in (container if container is not None else [])
                    if local_name(c.tag) in SME_TYPES
                ]
                index = siblings.index(current) if current in siblings else 0
                parts.append(f"[{index}]")
        current = parent_map.get(current)
    return "/".join(reversed(parts))


def fix_missing_property_values(root, logs):
    parent_map = build_parent_map(root)
    for prop in find_all_by_name(root, "property"):
        idshort = get_idshort(prop)
        inferred_type, inferred_value = infer_property_type_and_value(idshort)

        value_type = find_direct_child(prop, "valueType")
        if value_type is None:
            value_type = ET.Element(ns_tag(prop, "valueType"))
            insert_ordered(prop, value_type)
            value_type.text = inferred_type
            logs.append({
                "구분": "Property valueType",
                "요소": idshort,
                "수정 전": "없음",
                "수정 후": inferred_type,
                "결과": "Auto Fixed",
            })
        elif not text_of(value_type):
            value_type.text = inferred_type
            logs.append({
                "구분": "Property valueType",
                "요소": idshort,
                "수정 전": "빈 값",
                "수정 후": inferred_type,
                "결과": "Auto Fixed",
            })

        actual_type = text_of(value_type)

        value = find_direct_child(prop, "value")
        if value is None:
            value = ET.Element(ns_tag(prop, "value"))
            insert_ordered(prop, value)
            value.text = sample_value_for(actual_type, inferred_value)
            AUTO_VALUES.add(element_path(prop, parent_map))
            logs.append({
                "구분": "Property value 임의 입력 - 예시값 변경 필요",
                "요소": idshort or element_path(prop, parent_map),
                "수정 전": "없음",
                "수정 후": value.text,
                "결과": "Auto Fixed",
            })
        elif not text_of(value):
            value.text = sample_value_for(actual_type, inferred_value)
            AUTO_VALUES.add(element_path(prop, parent_map))
            logs.append({
                "구분": "Property value 임의 입력 - 예시값 변경 필요",
                "요소": idshort or element_path(prop, parent_map),
                "수정 전": "빈 값",
                "수정 후": value.text,
                "결과": "Auto Fixed",
            })


# =========================================================
# CD dataType <-> Property valueType 정렬
# =========================================================

MEASURE_DOWNGRADE = {
    "REAL_MEASURE": "REAL_COUNT",
    "INTEGER_MEASURE": "INTEGER_COUNT",
    "RATIONAL_MEASURE": "RATIONAL",
}

CD_DATATYPE_TO_XSD = {
    "STRING": "xs:string",
    "STRING_TRANSLATABLE": "xs:string",
    "IRDI": "xs:string",
    "IRI": "xs:anyURI",
    "HTML": "xs:string",
    "FILE": "xs:string",
    "BLOB": "xs:string",
    "RATIONAL": "xs:string",
    "RATIONAL_MEASURE": "xs:string",
    "REAL_MEASURE": "xs:float",
    "REAL_COUNT": "xs:float",
    "REAL_CURRENCY": "xs:float",
    "INTEGER_COUNT": "xs:integer",
    "INTEGER_MEASURE": "xs:integer",
    "INTEGER_CURRENCY": "xs:integer",
    "BOOLEAN": "xs:boolean",
    "DATE": "xs:date",
    "TIMESTAMP": "xs:dateTime",
    "TIME": "xs:time",
}


def compatible_cd_datatypes(property_value_type):
    value_type = normalize_value_type(property_value_type)

    # KOSMO Validator 가 실제로 보고한 "매칭 가능한 dataType" 목록에 맞춘다.
    if value_type in URI_TYPES:
        return {"IRI"}
    if value_type in STRING_TYPES:
        return {"STRING", "STRING_TRANSLATABLE", "IRDI", "HTML"}
    if value_type in REAL_TYPES:
        return {"REAL_MEASURE", "REAL_COUNT", "REAL_CURRENCY"}
    if value_type in INTEGER_TYPES:
        return {"INTEGER_COUNT", "INTEGER_MEASURE", "INTEGER_CURRENCY"}
    if value_type in BOOLEAN_TYPES:
        return {"BOOLEAN"}
    if value_type in {"xs:date", "date"}:
        return {"DATE"}
    if value_type in {"xs:dateTime", "dateTime"}:
        return {"TIMESTAMP"}
    if value_type in {"xs:time", "time"}:
        return {"TIME"}
    return None


def preferred_cd_datatype(property_value_type, unit_present):
    value_type = normalize_value_type(property_value_type)

    if value_type in URI_TYPES:
        return "IRI"
    if value_type in STRING_TYPES:
        return "STRING"
    if value_type in REAL_TYPES:
        return "REAL_MEASURE" if unit_present else "REAL_COUNT"
    if value_type in INTEGER_TYPES:
        return "INTEGER_MEASURE" if unit_present else "INTEGER_COUNT"
    if value_type in BOOLEAN_TYPES:
        return "BOOLEAN"
    if value_type in {"xs:date", "date"}:
        return "DATE"
    if value_type in {"xs:dateTime", "dateTime"}:
        return "TIMESTAMP"
    if value_type in {"xs:time", "time"}:
        return "TIME"
    return None


def get_iec61360_content(cd):
    for element in cd.iter():
        if local_name(element.tag) == "dataSpecificationIec61360":
            return element
    return None


def has_unit(cd):
    content = get_iec61360_content(cd)
    if content is None:
        return False
    if text_of(find_direct_child(content, "unit")):
        return True
    unit_id = find_direct_child(content, "unitId")
    return unit_id is not None and bool(get_reference_value(unit_id))


TYPED_ELEMENTS = ("property", "range")


def ensure_iec61360_content(cd):
    """IEC61360 content 가 없으면 골격을 만들어 반환한다."""
    content = get_iec61360_content(cd)
    if content is not None:
        return content
    idshort = get_idshort(cd)
    if not idshort:
        return None
    return build_iec61360(cd, idshort, None)


def align_property_and_cd_types(root, logs, refine_by_value=True):
    """
    Property/Range 의 valueType 과 CD 의 dataType 을 맞춘다.

    1) 선언된 valueType 이 구체적인데 실제 value 가 그 형식이 아니면
       value 를 근거로 valueType 을 다시 정한다. (예: xs:float + "1500 x 3000")
    2) valueType 이 xs:string(범용) 이고 CD 의 dataType 이 더 구체적이며
       value 가 그 형식에 맞으면 valueType 을 끌어올린다.
       (예: xs:string + REAL_MEASURE + "125.5" -> xs:float)
    3) 그 외에는 CD 의 dataType 을 valueType 에 맞춘다.
       CD 에 IEC61360 이나 dataType 이 없으면 새로 만든다.
    """
    cd_by_id = {
        get_id(cd): cd
        for cd in get_concept_descriptions(root)
        if get_id(cd)
    }

    for element_name in TYPED_ELEMENTS:
        for element in find_all_by_name(root, element_name):
            cd = cd_by_id.get(get_semantic_id_value(element))
            if cd is None:
                continue

            value_type_element = find_direct_child(element, "valueType")
            if value_type_element is None:
                continue

            property_type = text_of(value_type_element)
            if not property_type:
                continue

            idshort = get_idshort(element) or f"({element_name})"
            value_element = find_direct_child(element, "value")
            value_text = text_of(value_element)
            element_locked = is_protected_element(element)
            cd_locked = is_protected_cd(cd)

            content = get_iec61360_content(cd)
            datatype_element = (
                find_direct_child(content, "dataType")
                if content is not None else None
            )
            current_datatype = text_of(datatype_element)

            # --- 1) 값이 선언된 타입과 안 맞으면 값 기준으로 재추론 ---
            if (
                refine_by_value
                and not element_locked
                and value_text
                and not value_matches_type(value_text, property_type)
            ):
                inferred = infer_type_from_value(value_text)
                if inferred and inferred != property_type:
                    value_type_element.text = inferred
                    logs.append({
                        "구분": "Property valueType (값 기준 재추론)",
                        "요소": idshort,
                        "수정 전": f"{property_type} / value={value_text[:30]}",
                        "수정 후": inferred,
                        "결과": "Auto Fixed",
                    })
                    property_type = inferred

            compatible = compatible_cd_datatypes(property_type)
            if compatible is None:
                continue

            if current_datatype and current_datatype in compatible:
                continue

            # --- 2) xs:string 은 범용이므로 CD 쪽이 구체적이면 끌어올린다 ---
            target_type = CD_DATATYPE_TO_XSD.get(current_datatype)
            if cd_locked and element_locked:
                logs.append({
                    "구분": "valueType ↔ dataType (템플릿 원본 예외)",
                    "요소": idshort,
                    "수정 전": f"{property_type} vs CD {current_datatype or '없음'}",
                    "수정 후": "IDTA 원본 보호 - 수정 안 함",
                    "결과": "Warning",
                })
                continue
            if (
                not element_locked
                and normalize_value_type(property_type) in STRING_TYPES
                and target_type
                and normalize_value_type(target_type) not in STRING_TYPES
                and value_matches_type(value_text, target_type)
            ):
                value_type_element.text = target_type
                logs.append({
                    "구분": "Property valueType (CD 기준 상향)",
                    "요소": idshort,
                    "수정 전": f"{property_type} vs CD {current_datatype}",
                    "수정 후": target_type,
                    "결과": "Auto Fixed",
                })

                if value_element is not None and value_text:
                    normalized = normalize_value_text(value_text, target_type)
                    if normalized != value_text:
                        value_element.text = normalized
                        logs.append({
                            "구분": "Property value 표기 정규화",
                            "요소": idshort,
                            "수정 전": value_text,
                            "수정 후": normalized,
                            "결과": "Auto Fixed",
                        })
                continue

            # --- 3) CD 의 dataType 을 valueType 에 맞춘다 ---
            if cd_locked:
                logs.append({
                    "구분": "valueType ↔ dataType (표준 CD 보호)",
                    "요소": idshort,
                    "수정 전": f"{property_type} vs CD {current_datatype or '없음'}",
                    "수정 후": "표준 CD 보호 - 요소 valueType 또는 CD 수동 확인",
                    "결과": "Warning",
                })
                continue
            if content is None:
                content = ensure_iec61360_content(cd)
                if content is None:
                    logs.append({
                        "구분": "CD dataType",
                        "요소": get_idshort(cd) or "(idShort 없음)",
                        "수정 전": "IEC61360 없음",
                        "수정 후": "생성 불가 - 수동 확인 필요",
                        "결과": "Warning",
                    })
                    continue
                datatype_element = find_direct_child(content, "dataType")

            preferred = preferred_cd_datatype(property_type, has_unit(cd))
            if preferred is None:
                continue

            if datatype_element is None:
                datatype_element = ET.Element(ns_tag(content, "dataType"))
                insert_ordered(content, datatype_element)

            datatype_element.text = preferred
            logs.append({
                "구분": "CD dataType (valueType 기준 정렬)",
                "요소": get_idshort(cd),
                "수정 전": current_datatype or "없음",
                "수정 후": preferred,
                "결과": "Auto Fixed",
            })


# IEC61360 에서 존재하면 1자 이상이어야 하는 선택 요소
IEC61360_OPTIONAL_STRINGS = {
    "unit", "sourceOfDefinition", "symbol", "valueFormat", "value",
}

# 존재하면 1자 이상이어야 하는 일반 선택 요소
OPTIONAL_STRINGS = {
    "category", "contentType", "messageTopic", "version", "revision",
    "valueFormat", "symbol", "sourceOfDefinition", "unit",
    "assetType", "templateId", "lastUpdate", "minInterval", "maxInterval",
}

# 값이 비어 있으면 <value> 자체를 빼야 하는 요소 (File/Blob)
VALUE_OPTIONAL_PARENTS = {"file", "blob"}

# keys 가 비면 통째로 지워야 하는 선택 reference
OPTIONAL_REFERENCES = {
    "semanticId", "valueId", "unitId", "semanticIdListElement",
    "derivedFrom", "creator", "messageBroker", "externalSubjectId",
    "reference",
}

# 자식이 하나도 없으면 지워야 하는 컨테이너
EMPTY_CONTAINERS = {
    "description", "displayName", "preferredName", "shortName",
    "definition", "supplementalSemanticIds", "qualifiers",
    "embeddedDataSpecifications", "specificAssetIds", "extensions",
    "isCaseOf", "keys", "statements", "annotations", "valueList",
    "inputVariables", "outputVariables", "inoutputVariables",
}


def fix_empty_strings(root, logs):
    """
    IDTA 2-14 "String is shorter than 1 characters" 대응.

    - 값이 비어 있는 선택 문자열 요소를 제거한다.
    - text 가 빈 langString 항목을 제거한다.
    - 자식이 없는 컨테이너를 제거한다.
    """
    removed = 0

    for parent in list(root.iter()):
        parent_name = local_name(parent.tag)

        for child in list(parent):
            name = local_name(child.tag)

            # 1) 빈 선택 문자열
            if len(list(child)) == 0 and not text_of(child):
                in_iec = parent_name == "dataSpecificationIec61360"
                if (
                    (in_iec and name in IEC61360_OPTIONAL_STRINGS)
                    or (not in_iec and name in OPTIONAL_STRINGS)
                ):
                    parent.remove(child)
                    removed += 1
                    logs.append({
                        "구분": "빈 문자열 제거 (IDTA 2-14)",
                        "요소": get_idshort(parent) or parent_name,
                        "수정 전": f"<{name}/> (빈 값)",
                        "수정 후": "요소 제거",
                        "결과": "Auto Fixed",
                    })
                    continue

            # 1-1) File/Blob 의 빈 value -> <value> 제거 (contentType 만 유지)
            if (
                name == "value"
                and parent_name in VALUE_OPTIONAL_PARENTS
                and len(list(child)) == 0
                and not text_of(child)
            ):
                parent.remove(child)
                removed += 1
                logs.append({
                    "구분": "빈 File value 제거 (IDTA 2-14)",
                    "요소": get_idshort(parent) or parent_name,
                    "수정 전": "<value/>",
                    "수정 후": "value 요소 제거 (contentType 유지)",
                    "결과": "Auto Fixed",
                })
                continue

            # 1-2) path 가 빈 defaultThumbnail 제거
            if name == "defaultThumbnail" and not text_of(
                find_direct_child(child, "path")
            ):
                parent.remove(child)
                removed += 1
                continue

            # 2) text 가 빈 langString
            if name.startswith("langString"):
                text_element = find_direct_child(child, "text")
                if text_element is not None and not text_of(text_element):
                    parent.remove(child)
                    removed += 1
                    continue

    # 3) 남은 빈 컨테이너 / keys 없는 선택 reference 정리 (중첩 대비 3회)
    for _ in range(3):
        for parent in list(root.iter()):
            for child in list(parent):
                name = local_name(child.tag)
                if (
                    name in EMPTY_CONTAINERS
                    and len(list(child)) == 0
                    and not text_of(child)
                ):
                    parent.remove(child)
                    removed += 1
                    continue
                if name in OPTIONAL_REFERENCES and len(list(child)) > 0:
                    keys = find_direct_child(child, "keys")
                    if keys is None or not [
                        k for k in keys if local_name(k.tag) == "key"
                    ]:
                        parent.remove(child)
                        removed += 1

    if removed:
        logs.append({
            "구분": "빈 요소 정리 (IDTA 2-14)",
            "요소": "(전체)",
            "수정 전": "빈 문자열/컨테이너",
            "수정 후": f"{removed}건 제거",
            "결과": "Auto Fixed",
        })


def fix_measure_without_unit(root, logs):
    """AASc-3a-009: MEASURE 계열은 unit 또는 unitId 가 필수."""
    for cd in get_concept_descriptions(root):
        if is_protected_cd(cd):
            continue
        content = get_iec61360_content(cd)
        if content is None:
            continue

        datatype_element = find_direct_child(content, "dataType")
        current = text_of(datatype_element)
        if current not in MEASURE_DOWNGRADE:
            continue
        if has_unit(cd):
            continue

        datatype_element.text = MEASURE_DOWNGRADE[current]
        logs.append({
            "구분": "CD dataType (AASc-3a-009)",
            "요소": get_idshort(cd),
            "수정 전": f"{current} (unit 없음)",
            "수정 후": MEASURE_DOWNGRADE[current],
            "결과": "Auto Fixed",
        })


# =========================================================
# SubmodelElementList
# =========================================================

SME_TYPE_NAME = {
    "property": "Property",
    "multiLanguageProperty": "MultiLanguageProperty",
    "range": "Range",
    "file": "File",
    "blob": "Blob",
    "referenceElement": "ReferenceElement",
    "relationshipElement": "RelationshipElement",
    "annotatedRelationshipElement": "AnnotatedRelationshipElement",
    "entity": "Entity",
    "operation": "Operation",
    "capability": "Capability",
    "basicEventElement": "BasicEventElement",
    "submodelElementCollection": "SubmodelElementCollection",
    "submodelElementList": "SubmodelElementList",
}


def get_sml_direct_elements(sml):
    container = find_direct_child(sml, "value")
    if container is None:
        return []
    return [c for c in container if local_name(c.tag) in SME_TYPES]


def fix_sml_child_idshorts(root, logs):
    """
    AASd-120: (AAS 3.0) SML 의 직접 자식은 idShort 를 가질 수 없다.
    AAS 3.1 에서는 이 제약이 삭제되었으므로 기본(auto)은 3.0 일 때만 제거.
    """
    policy = POLICY.get("sml_item_idshort", "auto")
    if policy == "never":
        return
    if policy == "auto" and get_namespace(root.tag) != NS_AAS30:
        return
    for sml in find_all_by_name(root, "submodelElementList"):
        sml_idshort = get_idshort(sml)
        for child in get_sml_direct_elements(sml):
            idshort_element = find_direct_child(child, "idShort")
            if idshort_element is None:
                continue
            removed = idshort_element.text
            child.remove(idshort_element)
            logs.append({
                "구분": "SML 자식 idShort (AASd-120)",
                "요소": sml_idshort,
                "수정 전": removed or "",
                "수정 후": "제거",
                "결과": "Auto Fixed",
            })


def fix_sml_constraints(root, logs):
    for sml in find_all_by_name(root, "submodelElementList"):
        idshort = get_idshort(sml)
        children = get_sml_direct_elements(sml)

        child_types = [
            t for t in dict.fromkeys(
                SME_TYPE_NAME.get(local_name(c.tag)) for c in children
            ) if t
        ]

        type_element = find_direct_child(sml, "typeValueListElement")

        if len(child_types) == 1:
            expected = child_types[0]
            if type_element is None:
                type_element = ET.Element(ns_tag(sml, "typeValueListElement"))
                insert_ordered(sml, type_element)
                type_element.text = expected
                logs.append({
                    "구분": "SML typeValueListElement",
                    "요소": idshort,
                    "수정 전": "없음",
                    "수정 후": expected,
                    "결과": "Auto Fixed",
                })
            elif text_of(type_element) != expected:
                old = text_of(type_element)
                type_element.text = expected
                logs.append({
                    "구분": "SML typeValueListElement",
                    "요소": idshort,
                    "수정 전": old or "빈 값",
                    "수정 후": expected,
                    "결과": "Auto Fixed",
                })
        elif len(child_types) > 1:
            logs.append({
                "구분": "SML typeValueListElement (AASd-108)",
                "요소": idshort,
                "수정 전": " / ".join(child_types),
                "수정 후": "자식 타입 불일치 - 수동 확인 필요",
                "결과": "Warning",
            })
        elif type_element is None:
            logs.append({
                "구분": "SML typeValueListElement",
                "요소": idshort,
                "수정 전": "없음 (자식 없음)",
                "수정 후": "추론 불가 - 수동 입력 필요",
                "결과": "Warning",
            })

        if not children:
            continue

        if child_types and child_types[0] in {"Property", "Range"}:
            value_types = [
                v for v in dict.fromkeys(
                    text_of(find_direct_child(c, "valueType"))
                    for c in children
                ) if v
            ]
            if len(value_types) == 1:
                expected = value_types[0]
                vt_element = find_direct_child(sml, "valueTypeListElement")
                if vt_element is None:
                    vt_element = ET.Element(
                        ns_tag(sml, "valueTypeListElement")
                    )
                    insert_ordered(sml, vt_element)
                    vt_element.text = expected
                    logs.append({
                        "구분": "SML valueTypeListElement",
                        "요소": idshort,
                        "수정 전": "없음",
                        "수정 후": expected,
                        "결과": "Auto Fixed",
                    })
                elif text_of(vt_element) != expected:
                    old = text_of(vt_element)
                    vt_element.text = expected
                    logs.append({
                        "구분": "SML valueTypeListElement",
                        "요소": idshort,
                        "수정 전": old or "빈 값",
                        "수정 후": expected,
                        "결과": "Auto Fixed",
                    })
            elif len(value_types) > 1:
                logs.append({
                    "구분": "SML valueTypeListElement (AASd-109)",
                    "요소": idshort,
                    "수정 전": " / ".join(value_types),
                    "수정 후": "자식 valueType 불일치 - 수동 확인 필요",
                    "결과": "Warning",
                })

        semantic_values = [
            v for v in dict.fromkeys(
                get_semantic_id_value(c) for c in children
            ) if v
        ]

        list_element = find_direct_child(sml, "semanticIdListElement")
        if is_protected_element(sml) or any(
            is_standard_id(v) for v in semantic_values
        ):
            # IDTA 템플릿 SML 은 semanticIdListElement 구조를 바꾸지 않는다
            continue

        if len(semantic_values) > 1:
            if list_element is not None:
                sml.remove(list_element)
                logs.append({
                    "구분": "SML semanticIdListElement (AASd-107)",
                    "요소": idshort,
                    "수정 전": " / ".join(semantic_values[:3]) + " ...",
                    "수정 후": "자식 불일치 - semanticIdListElement 제거",
                    "결과": "Auto Fixed",
                })
            continue

        if len(semantic_values) == 1 and list_element is None:
            list_element = make_reference_element(
                sml, "semanticIdListElement", semantic_values[0]
            )
            insert_ordered(sml, list_element)
            logs.append({
                "구분": "SML semanticIdListElement 보충",
                "요소": idshort,
                "수정 전": "없음",
                "수정 후": semantic_values[0],
                "결과": "Auto Fixed",
            })
            continue

        if len(semantic_values) == 1 and list_element is not None:
            expected = semantic_values[0]
            old = get_reference_value(list_element)
            if old != expected:
                set_reference(list_element, expected)
                logs.append({
                    "구분": "SML semanticIdListElement",
                    "요소": idshort,
                    "수정 전": old or "없음",
                    "수정 후": expected,
                    "결과": "Auto Fixed",
                })


# =========================================================
# CD description / preferredName
# =========================================================

def has_english_lang_string(container):
    if container is None:
        return False
    for lang_string in container:
        language = text_of(find_direct_child(lang_string, "language"))
        text = text_of(find_direct_child(lang_string, "text"))
        if language.lower().startswith("en") and text:
            return True
    return False


SPECIAL_DESCRIPTIONS = {
    "EnergyEfficiency": "Information describing energy efficiency.",
    "CurrentPowerMeasurement":
        "Information describing the current power measurement.",
    "ActualData": "Information describing actual operating data.",
    "KPI": "Information describing a key performance indicator.",
}


def description_text_for(idshort):
    return f"{TBD_EN} Definition of {preferred_name_for(idshort)}"


def description_text_ko_for(idshort):
    return f"{TBD_KO} {idshort} 정의"


def preferred_name_ko_for(idshort):
    return f"{TBD_KO} {idshort}"


def preferred_name_for(idshort):
    tokens = split_tokens(idshort)
    if not tokens:
        return idshort
    return " ".join(t.capitalize() for t in tokens)


def add_lang_string(container, lang_string_tag, text_value, language="en"):
    lang_string = ET.SubElement(container, ns_tag(container, lang_string_tag))
    language_element = ET.SubElement(
        lang_string, ns_tag(lang_string, "language")
    )
    language_element.text = language
    text_element = ET.SubElement(lang_string, ns_tag(lang_string, "text"))
    text_element.text = text_value
    return lang_string


def build_iec61360(cd, name, datatype=None):
    container = find_direct_child(cd, "embeddedDataSpecifications")
    if container is None:
        container = ET.Element(ns_tag(cd, "embeddedDataSpecifications"))
        insert_ordered(cd, container)

    embedded = ET.SubElement(
        container, ns_tag(container, "embeddedDataSpecification")
    )

    data_spec = ET.SubElement(embedded, ns_tag(embedded, "dataSpecification"))
    set_reference(data_spec, IEC61360_TEMPLATE_ID)

    wrapper = ET.SubElement(
        embedded, ns_tag(embedded, "dataSpecificationContent")
    )
    content = ET.SubElement(
        wrapper, ns_tag(wrapper, "dataSpecificationIec61360")
    )

    preferred_name = ET.SubElement(content, ns_tag(content, "preferredName"))
    add_lang_string(
        preferred_name,
        "langStringPreferredNameTypeIec61360",
        preferred_name_for(name),
    )
    add_lang_string(
        preferred_name,
        "langStringPreferredNameTypeIec61360",
        preferred_name_ko_for(name),
        "ko",
    )

    if datatype:
        datatype_element = ET.SubElement(content, ns_tag(content, "dataType"))
        datatype_element.text = datatype

    definition = ET.SubElement(content, ns_tag(content, "definition"))
    add_lang_string(
        definition,
        "langStringDefinitionTypeIec61360",
        description_text_for(name),
    )
    add_lang_string(
        definition,
        "langStringDefinitionTypeIec61360",
        description_text_ko_for(name),
        "ko",
    )
    return content


def ensure_cd_description(cd, logs):
    description = find_direct_child(cd, "description")
    if has_english_lang_string(description):
        return

    idshort = get_idshort(cd)
    if not idshort:
        return

    if description is None:
        description = ET.Element(ns_tag(cd, "description"))
        insert_ordered(cd, description)

    text_value = description_text_for(idshort)
    add_lang_string(description, "langStringTextType", text_value)
    logs.append({
        "구분": "CD Description",
        "요소": idshort,
        "수정 전": "영문 description 없음",
        "수정 후": text_value,
        "결과": "Auto Fixed",
    })


def ensure_cd_preferred_name(cd, logs):
    content = get_iec61360_content(cd)
    if content is None:
        return

    preferred_name = find_direct_child(content, "preferredName")
    if has_english_lang_string(preferred_name):
        return

    idshort = get_idshort(cd)
    if preferred_name is None:
        preferred_name = ET.Element(ns_tag(content, "preferredName"))
        insert_ordered(content, preferred_name)

    text_value = preferred_name_for(idshort)
    add_lang_string(
        preferred_name, "langStringPreferredNameTypeIec61360", text_value
    )
    logs.append({
        "구분": "CD preferredName (IEC61360)",
        "요소": idshort,
        "수정 전": "없음",
        "수정 후": text_value,
        "결과": "Auto Fixed",
    })


def fix_concept_description_content(root, logs, create_iec61360=False):
    for cd in get_concept_descriptions(root):
        if create_iec61360 and get_iec61360_content(cd) is None:
            idshort = get_idshort(cd)
            if idshort:
                build_iec61360(cd, idshort, "STRING")
                logs.append({
                    "구분": "CD IEC61360 생성",
                    "요소": idshort,
                    "수정 전": "없음",
                    "수정 후": "preferredName / dataType=STRING",
                    "결과": "Auto Fixed",
                })
        ensure_cd_description(cd, logs)
        ensure_cd_preferred_name(cd, logs)


# =========================================================
# v3: AAS assetInformation (KOSMO-AAS #5, #6)
# =========================================================

def build_asset_url(idshort, version, revision):
    return (
        f"{BASE_URL}/asset/{remove_underscore(idshort)}/{version}/{revision}"
    )


ASSET_URL_PATTERN = re.compile(
    re.escape(BASE_URL) + r"/asset/[A-Za-z0-9\-]+/\d+(\.\d+)*/\d+(\.\d+)*$"
)


def fix_asset_information(root, logs):
    """assetKind / globalAssetId 정리. old->new globalAssetId 매핑 반환."""
    id_map = {}
    for shell in find_all_by_name(root, "assetAdministrationShell"):
        idshort = get_idshort(shell)
        info = find_direct_child(shell, "assetInformation")
        if info is None:
            info = ET.Element(ns_tag(shell, "assetInformation"))
            insert_ordered(shell, info)

        target_kind = POLICY.get("asset_kind")
        kind = find_direct_child(info, "assetKind")
        if target_kind and text_of(kind) != target_kind:
            old = text_of(kind)
            if kind is None:
                kind = ET.Element(ns_tag(info, "assetKind"))
                insert_ordered(info, kind)
            kind.text = target_kind
            logs.append({
                "구분": "assetKind (KOSMO-AAS #6)",
                "요소": idshort,
                "수정 전": old or "없음",
                "수정 후": target_kind,
                "결과": "Auto Fixed",
            })

        if not POLICY.get("fix_global_asset_id", True) or not idshort:
            continue
        version, revision = determine_version_revision(shell)
        expected = build_asset_url(idshort, version, revision)
        gaid = find_direct_child(info, "globalAssetId")
        old = text_of(gaid)
        # 3.0 초기 산출물처럼 globalAssetId 가 reference 형태면 값만 본다
        if gaid is not None and len(list(gaid)):
            old = get_reference_value(gaid) or ""
        if old == expected:
            continue
        if old and ASSET_URL_PATTERN.match(old) and old.split("/")[-3] == \
                remove_underscore(idshort):
            continue
        if gaid is None:
            gaid = ET.Element(ns_tag(info, "globalAssetId"))
            insert_ordered(info, gaid)
        for child in list(gaid):
            gaid.remove(child)
        gaid.text = expected
        if old:
            id_map[old] = expected
        logs.append({
            "구분": "globalAssetId (KOSMO-AAS #5)",
            "요소": idshort,
            "수정 전": old or "없음",
            "수정 후": expected,
            "결과": "Auto Fixed",
        })

    # Entity 등 다른 곳의 globalAssetId 도 같이 갱신
    if id_map:
        for element in find_all_by_name(root, "globalAssetId"):
            if text_of(element) in id_map:
                element.text = id_map[text_of(element)]
    return id_map


# =========================================================
# v3: 끊어진 ModelReference 복구
# =========================================================

SEMANTIC_REFERENCE_TAGS = {
    "semanticId", "semanticIdListElement", "dataSpecification", "unitId",
}


def fix_dangling_references(root, logs):
    """
    ModelReference 의 첫 key 가 Submodel 인데 그 id 를 가진 Submodel 이
    없으면, id 경로의 Submodel 이름(idShort)으로 찾아 복구한다.
    (AAS 이름을 바꾼 뒤 ReferenceElement 값이 옛 id 를 가리키는 경우)
    """
    submodels = find_all_by_name(root, "submodel")
    ids = {get_id(sm) for sm in submodels if get_id(sm)}
    by_idshort = {
        get_idshort(sm): get_id(sm) for sm in submodels if get_idshort(sm)
    }

    for reference in root.iter():
        keys = find_direct_child(reference, "keys")
        if keys is None:
            continue
        if text_of(find_direct_child(reference, "type")) != "ModelReference":
            continue
        key_list = [k for k in keys if local_name(k.tag) == "key"]
        if not key_list:
            continue
        first = key_list[0]
        if text_of(find_direct_child(first, "type")) != "Submodel":
            continue
        value_element = find_direct_child(first, "value")
        value = text_of(value_element)
        if not value or value in ids or is_standard_id(value):
            continue
        if local_name(reference.tag) in SEMANTIC_REFERENCE_TAGS:
            continue
        token = remove_underscore(get_last_id_token(value) or "")
        target = by_idshort.get(token)
        if target:
            value_element.text = target
            logs.append({
                "구분": "끊어진 참조 복구",
                "요소": local_name(reference.tag),
                "수정 전": value,
                "수정 후": target,
                "결과": "Auto Fixed",
            })
        else:
            logs.append({
                "구분": "끊어진 참조",
                "요소": local_name(reference.tag),
                "수정 전": value,
                "수정 후": "대응 Submodel 없음 - 수동 확인 필요",
                "결과": "Warning",
            })


# =========================================================
# v3: 특수 SME (KOSMO-SME #2) / 필수 Submodel (KOSMO-AAS #4)
# =========================================================

SPECIAL_SME = {"operation", "basicEventElement", "capability"}


def handle_special_sme(root, logs):
    remove = POLICY.get("remove_special_sme", False)
    for parent in list(root.iter()):
        for child in list(parent):
            name = local_name(child.tag)
            if name not in SPECIAL_SME:
                continue
            if remove:
                parent.remove(child)
            logs.append({
                "구분": "특수 SME (KOSMO-SME #2)",
                "요소": f"{SME_TYPE_NAME.get(name, name)} / {get_idshort(child)}",
                "수정 전": "참조모델에서 배제 대상",
                "수정 후": "제거" if remove else "수동 확인 필요",
                "결과": "Auto Fixed" if remove else "Warning",
            })


REQUIRED_SUBMODELS = {
    "DigitalNameplate": (
        lambda sem, ids: "nameplate" in sem.lower()
        or ids in {"Nameplate", "DigitalNameplate"}
    ),
    "HandoverDocumentation": (
        lambda sem, ids: "AHF578" in sem or ids == "HandoverDocumentation"
    ),
    "TechnicalData": (
        lambda sem, ids: "AHX837" in sem or ids == "TechnicalData"
    ),
    "OperationalData": (
        lambda sem, ids: "OperationalData" in sem or ids == "OperationalData"
    ),
}


def check_required_submodels(root):
    found = {name: False for name in REQUIRED_SUBMODELS}
    submodels = find_all_by_name(root, "submodel")
    for sm in submodels:
        sem = get_semantic_id_value(sm) or ""
        ids = get_idshort(sm)
        for name, test in REQUIRED_SUBMODELS.items():
            if test(sem, ids):
                found[name] = True
    missing = [n for n, ok in found.items() if not ok]
    return missing, len(submodels)


# =========================================================
# v3: CD 텍스트 (KOSMO-CD #3) / 한글 사전
# =========================================================

LANG_TAG = {
    "description": "langStringTextType",
    "displayName": "langStringNameType",
    "preferredName": "langStringPreferredNameTypeIec61360",
    "definition": "langStringDefinitionTypeIec61360",
}

# "Information about ..." 은 IDTA 원본 정의에도 쓰이므로 제외한다
PLACEHOLDER_PATTERN = re.compile(
    r"^(\[TBD\]|\[작성 필요\]|Information describing )", re.I
)


def lang_map(container):
    result = {}
    if container is None:
        return result
    for lang_string in container:
        language = text_of(find_direct_child(lang_string, "language")).lower()
        text = text_of(find_direct_child(lang_string, "text"))
        if language and text and language not in result:
            result[language] = text
    return result


def set_lang(parent, container_name, language, text, overwrite=False):
    """parent 의 container(langString 목록)에 language 텍스트를 넣는다."""
    container = find_direct_child(parent, container_name)
    if container is None:
        container = ET.Element(ns_tag(parent, container_name))
        insert_ordered(parent, container)
    for lang_string in container:
        if text_of(find_direct_child(lang_string, "language")).lower() == \
                language:
            current = text_of(find_direct_child(lang_string, "text"))
            if not overwrite and PLACEHOLDER_PATTERN.match(current) and \
                    not PLACEHOLDER_PATTERN.match(text):
                overwrite = True
            if overwrite and current != text:
                find_direct_child(lang_string, "text").text = text
                return True
            return False
    add_lang_string(container, LANG_TAG[container_name], text, language)
    return True


def apply_glossary_and_sync_cd_texts(root, logs):
    """
    자체(smart-factory) CD 에 대해
      1) 한글/영문 사전 값을 preferredName / definition / description /
         displayName 에 반영
      2) description <-> definition 서로 없는 언어를 복사
      3) displayName 이 없으면 preferredName 을 복사
    표준 CD 는 description 이 비었을 때 definition 복사만 한다.
    """
    overwrite = POLICY.get("glossary_overwrite", False)
    for cd in get_concept_descriptions(root):
        idshort = get_idshort(cd)
        content = get_iec61360_content(cd)
        protected = is_protected_cd(cd)
        changed = []

        entry = GLOSSARY_BY_ID.get(get_id(cd)) or GLOSSARY.get(idshort)
        if protected and not POLICY.get("standard_cd_korean", True):
            entry = None
        if entry and content is not None:
            # 표준 CD 영문은 비었거나 자리표시자일 때만 채운다 (원문 보호)
            pairs = [
                ("preferredName", "ko", entry.get("ko_name"), True),
                ("definition", "ko", entry.get("ko_definition"), True),
                ("preferredName", "en", entry.get("en_name"), not protected),
                ("definition", "en", entry.get("en_definition"),
                 not protected),
            ]
            # 검토 엑셀 값은 사용자가 고친 값이므로 덮어쓴다
            from_excel = overwrite or get_id(cd) in GLOSSARY_BY_ID
            for container, language, text, may_overwrite in pairs:
                force = from_excel and may_overwrite
                if text and set_lang(content, container, language, text,
                                     force):
                    changed.append(f"{container}@{language}")
                # description / displayName 도 같은 값으로 맞춘다
                mirror = {"definition": "description",
                          "preferredName": "displayName"}[container]
                # displayName 영문은 idShort(CamelCase) 규칙이라 맞추지 않는다
                if mirror == "displayName" and language == "en":
                    continue
                if text and set_lang(cd, mirror, language, text, force):
                    changed.append(f"{mirror}@{language}")

        definition = lang_map(
            find_direct_child(content, "definition")
            if content is not None else None
        )
        description = lang_map(find_direct_child(cd, "description"))

        copy_definition = definition if (not protected or not description) \
            else {}
        for language, text in copy_definition.items():
            if language not in description or (
                PLACEHOLDER_PATTERN.match(description.get(language, ""))
                and not PLACEHOLDER_PATTERN.match(text)
            ):
                if set_lang(cd, "description", language, text,
                            overwrite=language in description):
                    changed.append(f"description@{language}(definition 복사)")

        if not protected and content is not None:
            for language, text in description.items():
                if language not in definition:
                    if set_lang(content, "definition", language, text):
                        changed.append(f"definition@{language}")

            preferred = lang_map(find_direct_child(content, "preferredName"))
            for language, text in preferred.items():
                if set_lang(cd, "displayName", language, text):
                    changed.append(f"displayName@{language}")

        if changed:
            logs.append({
                "구분": "CD 텍스트 보충 (KOSMO-CD #3)",
                "요소": idshort,
                "수정 전": "누락",
                "수정 후": ", ".join(dict.fromkeys(changed)),
                "결과": "Auto Fixed",
            })


def fill_element_descriptions(root, logs):
    """자체 CD 를 쓰는 요소의 description 이 없으면 CD definition 을 복사."""
    cd_by_id = {get_id(cd): cd for cd in get_concept_descriptions(root)}
    count = 0
    for element in root.iter():
        if local_name(element.tag) not in SME_TYPES:
            continue
        semantic = get_semantic_id_value(element)
        if not is_self_id(semantic):
            continue
        cd = cd_by_id.get(semantic)
        content = get_iec61360_content(cd) if cd is not None else None
        definition = lang_map(
            find_direct_child(content, "definition")
            if content is not None else None
        )
        if not definition:
            continue
        existing = lang_map(find_direct_child(element, "description"))
        for language, text in definition.items():
            if language not in existing:
                set_lang(element, "description", language, text)
                count += 1
            elif PLACEHOLDER_PATTERN.match(existing[language]) and \
                    not PLACEHOLDER_PATTERN.match(text):
                set_lang(element, "description", language, text,
                         overwrite=True)
                count += 1
    if count:
        logs.append({
            "구분": "요소 description = CD definition",
            "요소": "(전체)",
            "수정 전": "description 누락",
            "수정 후": f"{count}건 CD definition 복사",
            "결과": "Auto Fixed",
        })


# =========================================================
# v3.1: CD 정의 생성 (LLM 등 외부 생성기) 지원
# =========================================================

# 마지막 실행에서 생성기가 채운 CD idShort (검토용 CSV 표시용)
GENERATED_CD = {}


def build_parent_map(root):
    return {child: parent for parent in root.iter() for child in parent}


def nearest_idshort_path(element, parent_map, limit=4):
    """요소의 상위 idShort 경로 (Submodel 까지, 최대 limit 단계)."""
    path = []
    current = parent_map.get(element)
    while current is not None and len(path) < limit:
        name = get_idshort(current)
        if name and local_name(current.tag) in SME_TYPES | {"submodel"}:
            path.append(name)
        if local_name(current.tag) == "submodel":
            break
        current = parent_map.get(current)
    return list(reversed(path))


def child_idshorts(element, limit=12):
    container = find_direct_child(element, "value")
    if container is None:
        container = find_direct_child(element, "statements")
    if container is None:
        return []
    names = []
    for child in container:
        if local_name(child.tag) not in SME_TYPES:
            continue
        name = get_idshort(child)
        if not name:
            # SML 항목: 항목 안의 idShort 를 대신 보여준다
            names += [n for n in child_idshorts(child, limit) if n]
        else:
            names.append(name)
        if len(names) >= limit:
            break
    return list(dict.fromkeys(names))[:limit]


def cd_needs_text(cd):
    """
    자체 CD: 이름/정의가 비었거나 자리표시자.
    표준 CD: 한글 이름/정의가 없음 (영문 원본은 유지, 옵션).
    """
    cd_id = get_id(cd)
    content = get_iec61360_content(cd)
    if not is_self_id(cd_id):
        if not is_standard_id(cd_id) or not POLICY.get(
            "standard_cd_korean", True
        ):
            return False
        if content is None:
            return False
        preferred = lang_map(find_direct_child(content, "preferredName"))
        definition = lang_map(find_direct_child(content, "definition"))
        return "ko" not in preferred or "ko" not in definition
    preferred = lang_map(find_direct_child(content, "preferredName")) \
        if content is not None else {}
    definition = lang_map(find_direct_child(content, "definition")) \
        if content is not None else {}
    description = lang_map(find_direct_child(cd, "description"))
    if not {"en", "ko"} <= set(preferred) or not {"en", "ko"} <= set(
        definition
    ):
        return True
    texts = list(definition.values()) + list(description.values())
    return any(PLACEHOLDER_PATTERN.match(t) for t in texts)


def collect_cd_text_targets(root, include_values=False):
    """
    정의 생성이 필요한 자체 CD 와 그 사용 맥락을 모은다.
    외부로 보낼 수 있도록 식별자(URL)는 넣지 않는다.
    """
    parent_map = build_parent_map(root)
    usage, list_usage = {}, {}
    for element in root.iter():
        if local_name(element.tag) not in SME_TYPES:
            continue
        semantic = get_semantic_id_value(element)
        if semantic and semantic not in usage:
            usage[semantic] = element
        list_element = find_direct_child(element, "semanticIdListElement")
        if list_element is not None:
            item_sem = get_reference_value(list_element)
            if item_sem:
                list_usage.setdefault(item_sem, element)

    targets = []
    for cd in get_concept_descriptions(root):
        idshort = get_idshort(cd)
        if not idshort or idshort in GLOSSARY or get_id(cd) in \
                GLOSSARY_BY_ID or not cd_needs_text(cd):
            continue
        content = get_iec61360_content(cd)
        definition = lang_map(find_direct_child(content, "definition")) \
            if content is not None else {}
        element = usage.get(get_id(cd))
        parent_list = list_usage.get(get_id(cd))
        target = {
            "idShort": idshort,
            "elementType": SME_TYPE_NAME.get(
                local_name(element.tag), local_name(element.tag)
            ) if element is not None else "",
            "valueType": text_of(find_direct_child(element, "valueType"))
            if element is not None else "",
            "unit": text_of(find_direct_child(content, "unit"))
            if content is not None else "",
            "path": nearest_idshort_path(element, parent_map)
            if element is not None else [],
            "children": child_idshorts(element)
            if element is not None else [],
        }
        if parent_list is not None:
            target["listItemOf"] = get_idshort(parent_list)
        if is_standard_id(get_id(cd)):
            # 표준 CD: 영문은 원본 그대로 쓰고 한글만 만든다
            preferred = lang_map(find_direct_child(content, "preferredName")) \
                if content is not None else {}
            target["standard"] = True
            target["given_en_name"] = preferred.get("en", "") or idshort
            target["given_en_definition"] = definition.get("en", "")
            target.pop("existing_en_definition", None)
        for language in ("en", "ko"):
            text = definition.get(language, "")
            if text and not PLACEHOLDER_PATTERN.match(text):
                target[f"existing_{language}_definition"] = text
        target = {k: v for k, v in target.items() if v not in ("", [], None)}
        if include_values and element is not None and \
                local_name(element.tag) == "property":
            target["exampleValue"] = text_of(
                find_direct_child(element, "value")
            )[:60]
        targets.append(target)
    return targets


def apply_generated_cd_texts(root, generated, logs, source="AI"):
    """
    생성된 이름/정의를 CD 에 반영한다.
    - 비어 있거나 자리표시자인 텍스트만 바꾼다 (사람이 쓴 정의는 보존)
    - preferredName ko / displayName 은 비었을 때만 채운다
    """
    applied = 0
    for cd in get_concept_descriptions(root):
        idshort = get_idshort(cd)
        entry = generated.get(idshort)
        if not entry:
            continue
        standard = is_standard_id(get_id(cd))
        if not is_self_id(get_id(cd)) and not standard:
            continue
        content = get_iec61360_content(cd) if standard \
            else ensure_iec61360_content(cd)
        if content is None:
            continue
        changed = []

        def put(parent, container, language, text):
            if not text:
                return
            current = lang_map(find_direct_child(parent, container))
            old = current.get(language)
            if old is None:
                set_lang(parent, container, language, text)
                changed.append(f"{container}@{language}")
            elif PLACEHOLDER_PATTERN.match(old) or (
                container == "preferredName"
                and language == "en"
                and old == preferred_name_for(idshort)
                and text != old
            ):
                set_lang(parent, container, language, text, overwrite=True)
                changed.append(f"{container}@{language}")

        if not standard:
            put(content, "preferredName", "en", entry.get("en_name"))
            put(content, "definition", "en", entry.get("en_definition"))
            put(cd, "description", "en", entry.get("en_definition"))
            put(cd, "displayName", "en", idshort)
        put(content, "preferredName", "ko", entry.get("ko_name"))
        put(content, "definition", "ko", entry.get("ko_definition"))
        put(cd, "description", "ko", entry.get("ko_definition"))
        # displayName 은 최종 preferredName 과 같게 맞춘다
        preferred = lang_map(find_direct_child(content, "preferredName"))
        put(cd, "displayName", "en", preferred.get("en") or idshort)
        put(cd, "displayName", "ko", preferred.get("ko") or entry.get(
            "ko_name"
        ))

        if changed:
            applied += 1
            GENERATED_CD[idshort] = source
            logs.append({
                "구분": (
                    f"CD 한글 생성 ({source}, IDTA 영문 유지) - 검토 필요"
                    if standard else f"CD 정의 생성 ({source}) - 검토 필요"
                ),
                "요소": idshort,
                "수정 전": "누락 / 자리표시자",
                "수정 후": f"{entry.get('ko_name', '')} | "
                          f"{entry.get('ko_definition', '')[:60]}",
                "결과": "Auto Fixed",
            })
    return applied


def export_cd_glossary_csv(xml_bytes):
    """
    자체 CD 의 현재 이름/정의를 사전 CSV 형식으로 내보낸다.
    검토 후 그대로 '한글/영문 사전 CSV' 로 다시 올릴 수 있다.
    """
    import csv
    root = ET.fromstring(xml_bytes)
    buffer = io.StringIO()
    fields = [
        "idShort", "ko_name", "ko_definition", "en_name", "en_definition",
        "구분", "출처", "검토", "상태",
    ]
    writer = csv.DictWriter(buffer, fieldnames=fields)
    writer.writeheader()
    rows = []
    for cd in get_concept_descriptions(root):
        standard = is_standard_id(get_id(cd))
        if not is_self_id(get_id(cd)) and not standard:
            continue
        idshort = get_idshort(cd)
        content = get_iec61360_content(cd)
        preferred = lang_map(find_direct_child(content, "preferredName")) \
            if content is not None else {}
        definition = lang_map(find_direct_child(content, "definition")) \
            if content is not None else {}
        source = GENERATED_CD.get(idshort) or (
            "사전" if idshort in GLOSSARY else "기존"
        )
        placeholder = any(
            PLACEHOLDER_PATTERN.match(t)
            for t in list(definition.values()) + list(preferred.values())
        )
        missing = not {"en", "ko"} <= set(definition)
        status = "AI 생성 - 검토" if source == "AI" else (
            "CD 자동 생성/자리표시자 - 작성 필요" if placeholder else (
                "한글 누락 - 작성 필요" if missing else ""
            )
        )
        rows.append({
            "idShort": idshort,
            "ko_name": preferred.get("ko", ""),
            "ko_definition": definition.get("ko", ""),
            "en_name": preferred.get("en", ""),
            "en_definition": definition.get("en", ""),
            "구분": "표준(IDTA·IRDI, 영문 원본 유지)" if standard else "자체",
            "출처": source,
            "검토": "필요" if status else "",
            "상태": status,
        })
    rows.sort(key=lambda r: (r["검토"] != "필요",))
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def report_cd_text_gaps(root, logs):
    """자체 CD 의 한글 누락 / 자리표시자 텍스트를 경고로 모아 보고."""
    no_ko, placeholder = [], []
    for cd in get_concept_descriptions(root):
        if is_protected_cd(cd):
            continue
        idshort = get_idshort(cd)
        content = get_iec61360_content(cd)
        preferred = lang_map(
            find_direct_child(content, "preferredName")
            if content is not None else None
        )
        definition = lang_map(
            find_direct_child(content, "definition")
            if content is not None else None
        )
        if "ko" not in preferred or "ko" not in definition:
            no_ko.append(idshort)
        texts = list(definition.values()) + list(
            lang_map(find_direct_child(cd, "description")).values()
        )
        if any(PLACEHOLDER_PATTERN.match(t) for t in texts):
            placeholder.append(idshort)
    if no_ko:
        logs.append({
            "구분": "CD 한글 누락 (KOSMO-CD #3)",
            "요소": f"{len(no_ko)}개 CD",
            "수정 전": ", ".join(no_ko[:15]) + (" ..." if len(no_ko) > 15 else ""),
            "수정 후": "한글 사전 CSV 로 보충 권장",
            "결과": "Warning",
        })
    if placeholder:
        logs.append({
            "구분": "CD 자리표시자 정의",
            "요소": f"{len(placeholder)}개 CD",
            "수정 전": ", ".join(placeholder[:15])
            + (" ..." if len(placeholder) > 15 else ""),
            "수정 후": "[작성 필요]/[TBD] 표시 → CD 정의 검토용 CSV 로 작성 후 재업로드",
            "결과": "Warning",
        })


# =========================================================
# 파이프라인
# =========================================================

def is_aas_environment(xml_bytes):
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return False
    return local_name(root.tag) == "environment"


def collect_namespaces(xml_bytes):
    namespace_map = {}
    try:
        for _, item in ET.iterparse(
            io.BytesIO(xml_bytes), events=("start-ns",)
        ):
            prefix, uri = item
            namespace_map.setdefault(prefix, uri)
    except ET.ParseError:
        pass
    return namespace_map


def edit_xml(xml_bytes, func):
    """네임스페이스 접두어를 보존하면서 XML 을 수정한다. func(root) -> result."""
    namespace_map = collect_namespaces(xml_bytes)
    backup = dict(getattr(ET, "_namespace_map", {}))
    try:
        for prefix, uri in namespace_map.items():
            try:
                ET.register_namespace(prefix, uri)
            except (ValueError, TypeError):
                pass
        root = ET.fromstring(xml_bytes)
        result = func(root)
        output = io.BytesIO()
        ET.ElementTree(root).write(
            output, encoding="utf-8", xml_declaration=True
        )
        return output.getvalue(), result
    finally:
        ET._namespace_map.clear()
        ET._namespace_map.update(backup)


def apply_policy(options):
    for key in POLICY:
        if key in options:
            POLICY[key] = options[key]
    SUBMODEL_KIND["value"] = options.get("submodel_kind", "Template")
    GLOSSARY.clear()
    GLOSSARY.update(options.get("glossary") or {})
    GLOSSARY_BY_ID.clear()
    GLOSSARY_BY_ID.update(options.get("glossary_by_id") or {})


def run_rules(root, xml_path, options, logs):
    # AAS 3.0 idShort 는 하이픈 불가, 3.1 은 허용
    POLICY["allow_hyphen"] = get_namespace(root.tag) != NS_AAS30
    # 1. idShort / displayName
    clean_idshorts(root, logs)
    clean_display_names(root, logs)

    # 2. AAS administration + ID + assetInformation
    aas_map = fix_aas_administration(root, logs)
    remap_all_keys(root, aas_map, logs, "AAS ID")
    fix_asset_information(root, logs)

    # 3. Submodel (ID 재작성 전에 소유 관계 캡처)
    owner_map = build_submodel_owner_map(root)
    default_equipment = get_default_equipment_name(root, xml_path)
    sm_map, idshort_map = fix_submodels(
        root, owner_map, default_equipment, logs
    )
    remap_all_keys(root, sm_map, logs, "Submodel ID")
    normalize_aas_submodel_references(root, idshort_map, logs)
    fix_dangling_references(root, logs)

    # 4. ConceptDescription ID (표준 id 는 보호) + 참조 리맵
    fix_cd_missing_idshort(root, logs)
    cd_map = normalize_concept_descriptions(root, logs)
    remap_all_keys(root, cd_map, logs, "ConceptDescription ID")
    remove_duplicate_concept_descriptions(root, logs)

    # 5. 특수 SME
    handle_special_sme(root, logs)

    # 6. semanticId 해소 (템플릿 IRDI/IRI 처리 포함)
    resolve_semantic_ids(root, logs, options.get("create_missing_cd", True))

    # 7. Property valueType 세분화 -> 실수 통일 -> 값 보충 -> CSV 예시값
    AUTO_VALUES.clear()
    if options.get("refine_value_types", True):
        refine_property_value_types(root, logs)
    unify_real_value_types(root, logs)
    fix_missing_property_values(root, logs)
    apply_value_overrides(root, options.get("value_overrides") or {}, logs)

    # 8. 타입 정렬
    align_property_and_cd_types(
        root, logs, options.get("refine_by_value", True)
    )
    fix_measure_without_unit(root, logs)

    # 9. SubmodelElementList
    fix_sml_child_idshorts(root, logs)
    fix_sml_constraints(root, logs)

    # 10. CD 텍스트 (사전 반영 → AI 생성 → 동기화 → 자리표시자 → 요소 description)
    GENERATED_CD.clear()
    apply_glossary_and_sync_cd_texts(root, logs)
    generator = options.get("cd_text_generator")
    if generator is not None:
        targets = collect_cd_text_targets(
            root, options.get("llm_send_values", False)
        )
        if targets:
            generated, errors = generator(targets)
            apply_generated_cd_texts(root, generated, logs)
            for error in errors:
                logs.append({
                    "구분": "CD 정의 생성 실패",
                    "요소": error.get("idShort", "(배치)"),
                    "수정 전": "생성 요청",
                    "수정 후": error.get("message", "")[:200],
                    "결과": "Warning",
                })
            apply_glossary_and_sync_cd_texts(root, logs)
    fix_concept_description_content(
        root, logs, options.get("create_iec61360", False)
    )
    if options.get("fill_element_description", True):
        fill_element_descriptions(root, logs)
    report_cd_text_gaps(root, logs)

    # 11. 빈 문자열 / 빈 컨테이너 (IDTA 2-14)
    fix_empty_strings(root, logs)

    # 12. 요소 순서 (XSD sequence)
    if options.get("reorder", True):
        reorder_all(root, logs)

    return [
        remove_underscore(get_idshort(s))
        for s in find_all_by_name(root, "assetAdministrationShell")
        if get_idshort(s)
    ]


def process_aas_xml(xml_bytes, xml_path, options):
    apply_policy(options)
    logs = []
    fixed, names = edit_xml(
        xml_bytes, lambda root: run_rules(root, xml_path, options, logs)
    )
    return fixed, logs, names


# =========================================================
# AASX 패키지 (KOSMO-AAS #1, D-05)
# =========================================================

REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_SUPPL = "http://admin-shell.io/aasx/relationships/aas-suppl"
REL_SPEC = "http://admin-shell.io/aasx/relationships/aas-spec"
REL_THUMB = (
    "http://schemas.openxmlformats.org/package/2006/relationships/"
    "metadata/thumbnail"
)

MIME_BY_EXT = {
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
    "gif": "image/gif", "bmp": "image/bmp", "svg": "image/svg+xml",
    "webp": "image/webp", "pdf": "application/pdf", "xml": "text/xml",
    "json": "application/json", "txt": "text/plain", "csv": "text/csv",
    "rels": "application/vnd.openxmlformats-package.relationships+xml",
    "zip": "application/zip", "step": "application/step",
    "stp": "application/step", "html": "text/html", "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet",
}


def part_name(path):
    return "/" + path.lstrip("/")


def rels_path_for(part):
    part = part.lstrip("/")
    directory, _, name = part.rpartition("/")
    return f"{directory + '/' if directory else ''}_rels/{name}.rels"


def resolve_target(source_part, target):
    if target.startswith("/"):
        return target
    base = part_name(source_part).rsplit("/", 1)[0]
    parts = (base + "/" + target).split("/")
    stack = []
    for piece in parts:
        if piece in ("", "."):
            continue
        if piece == "..":
            if stack:
                stack.pop()
        else:
            stack.append(piece)
    return "/" + "/".join(stack)


def read_rels(data):
    if not data:
        return []
    root = ET.fromstring(data)
    return [
        dict(rel.attrib) for rel in root
        if local_name(rel.tag) == "Relationship"
    ]


def write_rels(relations):
    lines = [
        '<?xml version="1.0" encoding="utf-8"?>',
        f'<Relationships xmlns="{REL_NS}">',
    ]
    for rel in relations:
        attrs = " ".join(
            f'{k}="{v}"' for k, v in rel.items()
            if k in ("Type", "Target", "Id", "TargetMode")
        )
        lines.append(f"  <Relationship {attrs} />")
    lines.append("</Relationships>")
    return "\n".join(lines).encode("utf-8")


def next_rel_id(relations):
    used = {r.get("Id") for r in relations}
    index = 1
    while f"R{index}" in used:
        index += 1
    return f"R{index}"


def collect_internal_file_refs(xml_bytes):
    """XML 안 File value / defaultThumbnail path 중 패키지 내부 경로."""
    root = ET.fromstring(xml_bytes)
    refs = set()
    for element in root.iter():
        name = local_name(element.tag)
        if name == "file":
            value = text_of(find_direct_child(element, "value"))
            if value.startswith("/") and not value.startswith("//"):
                refs.add(value)
        if name == "defaultThumbnail":
            path = text_of(find_direct_child(element, "path"))
            if path.startswith("/"):
                refs.add(path)
    return refs


def thumbnail_paths(xml_bytes):
    root = ET.fromstring(xml_bytes)
    return [
        text_of(find_direct_child(t, "path"))
        for t in find_all_by_name(root, "defaultThumbnail")
        if text_of(find_direct_child(t, "path"))
    ]


def set_default_thumbnail(xml_bytes, path):
    mime = MIME_BY_EXT.get(path.rsplit(".", 1)[-1].lower(), "image/png")

    def apply(root):
        for shell in find_all_by_name(root, "assetAdministrationShell"):
            info = find_direct_child(shell, "assetInformation")
            if info is None or find_direct_child(info, "defaultThumbnail") \
                    is not None:
                continue
            thumb = ET.Element(ns_tag(info, "defaultThumbnail"))
            p = ET.SubElement(thumb, ns_tag(thumb, "path"))
            p.text = path
            c = ET.SubElement(thumb, ns_tag(thumb, "contentType"))
            c.text = mime
            insert_ordered(info, thumb)
            return True
        return False

    return edit_xml(xml_bytes, apply)


def fix_package(entries, env_parts, options, logs):
    """
    entries: {name: bytes} (디렉터리 엔트리 제외), env_parts: XML 파트 목록.
    결과를 entries 에 반영하고 패키지 점검 정보를 반환한다.
    """
    info = {"thumbnail": None, "missing_files": [], "unreferenced": []}

    # --- 1) 파일명 언더바 제거 (선택) ---
    if options.get("rename_underscore_files", True):
        all_refs = set()
        for part in env_parts:
            all_refs |= collect_internal_file_refs(entries[part])
        for ref in sorted(all_refs):
            directory, _, base = ref.rpartition("/")
            if "_" not in base or ref.lstrip("/") not in entries:
                continue
            new_base = base.replace("_", "")
            new_ref = f"{directory}/{new_base}"
            while new_ref.lstrip("/") in entries:
                stem, _, ext = new_base.rpartition(".")
                new_base = f"{stem}1.{ext}" if stem else new_base + "1"
                new_ref = f"{directory}/{new_base}"
            entries[new_ref.lstrip("/")] = entries.pop(ref.lstrip("/"))
            for part in env_parts:
                entries[part] = entries[part].replace(
                    f">{ref}<".encode(), f">{new_ref}<".encode()
                )
            for name in list(entries):
                if name.endswith(".rels"):
                    entries[name] = entries[name].replace(
                        f'"{ref}"'.encode(), f'"{new_ref}"'.encode()
                    )
            logs.append({
                "구분": "파일명 언더바 제거 (D-05)",
                "요소": base,
                "수정 전": ref,
                "수정 후": new_ref,
                "결과": "Auto Fixed",
            })

    # --- 2) 참조 파일 존재 / aas-suppl relationship ---
    referenced_all = set()
    for part in env_parts:
        refs = collect_internal_file_refs(entries[part])
        thumbs = set(thumbnail_paths(entries[part]))
        referenced_all |= refs
        for ref in sorted(refs):
            if ref.lstrip("/") not in entries:
                info["missing_files"].append(ref)
                logs.append({
                    "구분": "File 참조 파일 없음 (D-05)",
                    "요소": part,
                    "수정 전": ref,
                    "수정 후": "패키지에 파일 추가 또는 value 제거 필요",
                    "결과": "Warning",
                })

        rels_name = rels_path_for(part)
        relations = read_rels(entries.get(rels_name))
        kept, removed = [], []
        for rel in relations:
            if rel.get("Type") != REL_SUPPL:
                kept.append(rel)
                continue
            target = resolve_target(part, rel.get("Target", ""))
            if target.lstrip("/") not in entries:
                removed.append(target)
                continue
            if options.get("remove_unreferenced", True) and target not in refs:
                removed.append(target)
                continue
            kept.append(rel)
        existing = {
            resolve_target(part, r.get("Target", ""))
            for r in kept if r.get("Type") == REL_SUPPL
        }
        added = []
        for ref in sorted(refs - thumbs):
            if ref in existing or ref.lstrip("/") not in entries:
                continue
            kept.append({
                "Type": REL_SUPPL, "Target": ref, "Id": next_rel_id(kept),
            })
            added.append(ref)
        if removed or added or rels_name not in entries:
            entries[rels_name] = write_rels(kept)
        if added:
            logs.append({
                "구분": "aas-suppl relationship 추가 (D-05)",
                "요소": part,
                "수정 전": "누락",
                "수정 후": ", ".join(added),
                "결과": "Auto Fixed",
            })
        if removed:
            logs.append({
                "구분": "aas-suppl relationship 정리 (D-05)",
                "요소": part,
                "수정 전": ", ".join(removed),
                "수정 후": "참조 없음/파일 없음 → 제거",
                "결과": "Auto Fixed",
            })

    # --- 3) 참조되지 않는 보조 파일 제거 (선택) ---
    if options.get("remove_unreferenced", True):
        thumb_targets = set()
        for rel in read_rels(entries.get("_rels/.rels")):
            if rel.get("Type") == REL_THUMB:
                thumb_targets.add(resolve_target("", rel.get("Target", "")))
        for part in env_parts:
            thumb_targets |= set(thumbnail_paths(entries[part]))
        for name in list(entries):
            path = part_name(name)
            if not name.startswith("aasx/") or name.endswith(".rels"):
                continue
            if name in env_parts or name == "aasx/aasx-origin":
                continue
            if path in referenced_all or path in thumb_targets:
                continue
            del entries[name]
            info["unreferenced"].append(path)
        if info["unreferenced"]:
            logs.append({
                "구분": "참조되지 않는 파일 제거 (D-05)",
                "요소": "(패키지)",
                "수정 전": ", ".join(info["unreferenced"]),
                "수정 후": "제거",
                "결과": "Auto Fixed",
            })

    # --- 4) 썸네일 (KOSMO-AAS #1) ---
    root_rels = read_rels(entries.get("_rels/.rels"))
    rel_thumb = [
        resolve_target("", r.get("Target", ""))
        for r in root_rels if r.get("Type") == REL_THUMB
    ]
    rel_thumb = [t for t in rel_thumb if t.lstrip("/") in entries]
    xml_thumb = []
    for part in env_parts:
        xml_thumb += [
            t for t in thumbnail_paths(entries[part])
            if t.lstrip("/") in entries
        ]
    if xml_thumb:
        info["thumbnail"] = xml_thumb[0]
        if xml_thumb[0] not in rel_thumb:
            root_rels = [r for r in root_rels if r.get("Type") != REL_THUMB]
            root_rels.append({
                "Type": REL_THUMB, "Target": xml_thumb[0],
                "Id": next_rel_id(root_rels),
            })
            entries["_rels/.rels"] = write_rels(root_rels)
            logs.append({
                "구분": "썸네일 relationship (KOSMO-AAS #1)",
                "요소": "_rels/.rels",
                "수정 전": "없음/불일치",
                "수정 후": xml_thumb[0],
                "결과": "Auto Fixed",
            })
    elif rel_thumb and env_parts:
        new_xml, done = set_default_thumbnail(entries[env_parts[0]],
                                              rel_thumb[0])
        if done:
            entries[env_parts[0]] = new_xml
            info["thumbnail"] = rel_thumb[0]
            logs.append({
                "구분": "defaultThumbnail (KOSMO-AAS #1)",
                "요소": "assetInformation",
                "수정 전": "없음",
                "수정 후": rel_thumb[0],
                "결과": "Auto Fixed",
            })
    else:
        logs.append({
            "구분": "썸네일 없음 (KOSMO-AAS #1)",
            "요소": "(패키지)",
            "수정 전": "썸네일 이미지 없음",
            "수정 후": "장비 대표 이미지(비식별) 추가 필요",
            "결과": "Warning",
        })

    # --- 5) [Content_Types].xml ---
    ct_name = "[Content_Types].xml"
    ct_data = entries.get(ct_name)
    defaults, overrides = {}, []
    if ct_data:
        ct_root = ET.fromstring(ct_data)
        for child in ct_root:
            if local_name(child.tag) == "Default":
                defaults[child.get("Extension", "").lower()] = child.get(
                    "ContentType"
                )
            elif local_name(child.tag) == "Override":
                overrides.append(
                    (child.get("PartName"), child.get("ContentType"))
                )
    added = []
    for name in entries:
        if "." not in name.rsplit("/", 1)[-1] or name == ct_name:
            continue
        ext = name.rsplit(".", 1)[-1].lower()
        overridden = any(p == part_name(name) for p, _ in overrides)
        if ext not in defaults and not overridden:
            defaults[ext] = MIME_BY_EXT.get(ext, "application/octet-stream")
            added.append(ext)
    if added or not ct_data:
        lines = [
            '<?xml version="1.0" encoding="utf-8"?>',
            f'<Types xmlns="{CT_NS}">',
        ]
        for ext, mime in defaults.items():
            lines.append(f'  <Default Extension="{ext}" ContentType="{mime}" />')
        for part, mime in overrides:
            lines.append(f'  <Override PartName="{part}" ContentType="{mime}" />')
        lines.append("</Types>")
        entries[ct_name] = "\n".join(lines).encode("utf-8")
        logs.append({
            "구분": "Content_Types 등록 (D-05)",
            "요소": ct_name,
            "수정 전": "확장자 누락",
            "수정 후": ", ".join(added) or "(신규 생성)",
            "결과": "Auto Fixed",
        })

    return info


def process_aasx(file_bytes, options):
    logs = []
    processed_files, skipped_json, equipment_names = [], [], []
    order, entries = [], {}

    with zipfile.ZipFile(io.BytesIO(file_bytes), "r") as input_zip:
        for item in input_zip.infolist():
            if item.is_dir() or item.filename.endswith("/"):
                logs.append({
                    "구분": "ZIP 디렉터리 엔트리 제거",
                    "요소": item.filename,
                    "수정 전": "디렉터리 엔트리",
                    "수정 후": "제거",
                    "결과": "Auto Fixed",
                })
                continue
            order.append(item.filename)
            entries[item.filename] = input_zip.read(item.filename)

    env_parts = []
    for name in order:
        data = entries[name]
        lower = name.lower()
        if lower.endswith(".json") and b'"assetAdministrationShells"' in data:
            skipped_json.append(name)
            continue
        if lower.endswith(".xml") and is_aas_environment(data):
            fixed, file_logs, names = process_aas_xml(data, name, options)
            entries[name] = fixed
            logs.extend(file_logs)
            processed_files.append(name)
            equipment_names.extend(names)
            env_parts.append(name)

    package_info = {}
    if env_parts and options.get("fix_package", True):
        package_info = fix_package(entries, env_parts, options, logs)

    output_buffer = io.BytesIO()
    with zipfile.ZipFile(output_buffer, "w", zipfile.ZIP_DEFLATED) as out:
        names = ["[Content_Types].xml"] if "[Content_Types].xml" in entries \
            else []
        names += [n for n in order if n in entries and n not in names]
        names += [n for n in entries if n not in names]
        for name in names:
            out.writestr(name, entries[name])

    return {
        "bytes": output_buffer.getvalue(),
        "logs": logs,
        "processed_files": processed_files,
        "skipped_json": skipped_json,
        "equipment_names": list(dict.fromkeys(equipment_names)),
        "xml_texts": {p: entries[p] for p in env_parts},
        "package_info": package_info,
        "entries": entries,
    }


# =========================================================
# Validator 항목 기준 점검 리포트
# =========================================================

CD_URL_PATTERN = re.compile(
    re.escape(BASE_URL) + r"/cd/([A-Za-z0-9\-]+)/\d+(\.\d+)*/\d+(\.\d+)*$"
)
SM_URL_PATTERN = re.compile(
    re.escape(BASE_URL)
    + r"/sm/([A-Za-z0-9\-]+)/([A-Za-z0-9\-]+)/\d+(\.\d+)*/\d+(\.\d+)*$"
)


def _row(item, problems, warn=False, note=""):
    if problems and all("[표준CD]" in str(p) for p in problems):
        warn = True
        note = "IDTA 템플릿 원본 CD 예외 (가이드라인 7: 수정하지 않음)"
    if not problems:
        status = "✅ 통과"
    else:
        status = "⚠️ 확인" if warn else "❌ 오류"
    detail = "; ".join(str(p) for p in problems[:8])
    if len(problems) > 8:
        detail += f" ... 외 {len(problems) - 8}건"
    if note and problems:
        detail = f"{note} - {detail}"
    return {
        "검사항목": item,
        "결과": status,
        "건수": len(problems),
        "상세": detail or note,
    }


def build_report(xml_bytes, package_info=None):
    """추진단 Validator 항목 번호 기준으로 결과를 다시 점검한다."""
    root = ET.fromstring(xml_bytes)
    rows = []
    cds = get_concept_descriptions(root)
    cd_by_id = {get_id(cd): cd for cd in cds if get_id(cd)}
    shells = find_all_by_name(root, "assetAdministrationShell")
    submodels = find_all_by_name(root, "submodel")

    # ---- KOSMO-AAS ----
    if package_info is not None:
        rows.append(_row(
            "KOSMO-AAS #1 Thumbnail",
            [] if package_info.get("thumbnail") else ["썸네일 없음"],
        ))
    rows.append(_row("KOSMO-AAS #2 idShort", [
        get_idshort(s) for s in shells
        if not IDSHORT_PATTERN.match(get_idshort(s) or "")
    ]))
    rows.append(_row("KOSMO-AAS #3 Id 형식", [
        f"{get_id(s)}" for s in shells
        if get_id(s) != build_aas_url(
            get_idshort(s), *determine_version_revision(s)
        )
    ]))
    missing, count = check_required_submodels(root)
    problems = [f"{m} 없음" for m in missing]
    if count < 6:
        problems.append(f"Submodel {count}종 (6종 이상 권장)")
    rows.append(_row("KOSMO-AAS #4 필수 Submodel", problems))
    gaid_problems = []
    for s in shells:
        info = find_direct_child(s, "assetInformation")
        gaid = text_of(find_direct_child(info, "globalAssetId"))
        if not ASSET_URL_PATTERN.match(gaid or ""):
            gaid_problems.append(gaid or "없음")
    rows.append(_row("KOSMO-AAS #5 globalAssetId", gaid_problems))
    rows.append(_row("KOSMO-AAS #6 Type", [
        text_of(find_direct_child(
            find_direct_child(s, "assetInformation"), "assetKind"
        )) or "없음"
        for s in shells
        if text_of(find_direct_child(
            find_direct_child(s, "assetInformation"), "assetKind"
        )) != "Type"
    ]))

    # ---- KOSMO-Submodel ----
    aas_names = {remove_underscore(get_idshort(s)) for s in shells}
    rows.append(_row("KOSMO-Submodel #1 idShort", [
        get_idshort(sm) for sm in submodels
        if not IDSHORT_PATTERN.match(get_idshort(sm) or "")
    ]))
    sm_id_problems = []
    for sm in submodels:
        m = SM_URL_PATTERN.match(get_id(sm) or "")
        if not m or m.group(2) != get_idshort(sm) or (
            aas_names and m.group(1) not in aas_names
        ):
            sm_id_problems.append(f"{get_idshort(sm)}: {get_id(sm)}")
    rows.append(_row("KOSMO-Submodel #2 Id 형식", sm_id_problems))
    sem_problems = []
    for sm in submodels:
        sem = get_semantic_id_value(sm)
        if not sem:
            sem_problems.append(f"{get_idshort(sm)}: 없음")
        elif not (is_standard_id(sem) or is_self_id(sem)):
            sem_problems.append(f"{get_idshort(sm)}: {sem}")
    rows.append(_row("KOSMO-Submodel #3 semanticId", sem_problems))
    rows.append(_row("KOSMO-Submodel #4 Kind", [
        f"{get_idshort(sm)}: {text_of(find_direct_child(sm, 'kind')) or '없음'}"
        for sm in submodels
        if text_of(find_direct_child(sm, "kind")) != "Template"
    ]))

    # ---- KOSMO-SME ----
    sme_idshort, special, no_cd, no_value, type_mismatch = [], [], [], [], []
    for element in root.iter():
        name = local_name(element.tag)
        if name not in SME_TYPES:
            continue
        idshort = get_idshort(element)
        if idshort and not IDSHORT_PATTERN.match(idshort):
            sme_idshort.append(idshort)
        if name in SPECIAL_SME:
            special.append(f"{SME_TYPE_NAME[name]}/{idshort}")
        sem = get_semantic_id_value(element)
        if not sem or sem not in cd_by_id:
            no_cd.append(f"{idshort or name}: {sem or '없음'}")
        if name == "property" and not text_of(
            find_direct_child(element, "value")
        ):
            no_value.append(idshort or "(SML 항목)")
        if name in TYPED_ELEMENTS and sem in cd_by_id:
            value_type = text_of(find_direct_child(element, "valueType"))
            compatible = compatible_cd_datatypes(value_type)
            content = get_iec61360_content(cd_by_id[sem])
            datatype = text_of(find_direct_child(content, "dataType")) \
                if content is not None else ""
            if compatible and datatype not in compatible:
                tag = " [표준CD]" if is_standard_id(sem) else ""
                type_mismatch.append(
                    f"{idshort}: {value_type} vs {datatype or '없음'}{tag}"
                )
    rows.append(_row("KOSMO-SME #1 idShort", sme_idshort))
    rows.append(_row("KOSMO-SME #2 특수 SME 배제", special))
    rows.append(_row("KOSMO-SME #3 CD 매핑", no_cd))
    rows.append(_row("KOSMO-SME #4 Property value", no_value))
    rows.append(_row("KOSMO-SME #5 valueType = CD dataType", type_mismatch))

    # ---- KOSMO-CD ----
    used_typed = set()
    for element_name in TYPED_ELEMENTS:
        for element in find_all_by_name(root, element_name):
            used_typed.add(get_semantic_id_value(element))
    cd_idshort, cd_id, cd_text, cd_dtype, seen, dups = [], [], [], [], set(), []
    for cd in cds:
        idshort, cid = get_idshort(cd), get_id(cd)
        if not IDSHORT_PATTERN.match(idshort or ""):
            cd_idshort.append(idshort or "(없음)")
        if is_self_id(cid):
            m = CD_URL_PATTERN.match(cid)
            if not m or m.group(1) != idshort:
                cd_id.append(f"{idshort}: {cid}")
        elif not is_standard_id(cid):
            cd_id.append(f"{idshort}: {cid}")
        content = get_iec61360_content(cd)
        definition = lang_map(find_direct_child(content, "definition")) \
            if content is not None else {}
        description = lang_map(find_direct_child(cd, "description"))
        if not definition and not description:
            cd_text.append(idshort)
        if cid in used_typed:
            if content is None or not text_of(
                find_direct_child(content, "dataType")
            ):
                tag = " [표준CD]" if is_standard_id(cid) else ""
                cd_dtype.append(f"{idshort}{tag}")
        key = (idshort, cid)
        if key in seen:
            dups.append(f"{idshort}: {cid}")
        seen.add(key)
    rows.append(_row("KOSMO-CD #1 idShort", cd_idshort))
    rows.append(_row("KOSMO-CD #2 Id 형식", cd_id))
    rows.append(_row("KOSMO-CD #3 definition/description", cd_text))
    rows.append(_row("KOSMO-CD #4 DataType", cd_dtype))
    rows.append(_row("KOSMO-CD #5 idShort/id 중복", dups))

    # ---- 작성 필요 표시 / 임의 예시값 ----
    tbd = []
    for cd in cds:
        content = get_iec61360_content(cd)
        texts = list(lang_map(find_direct_child(cd, "description")).values())
        if content is not None:
            texts += list(lang_map(find_direct_child(content, "definition"))
                          .values())
            texts += list(lang_map(find_direct_child(content, "preferredName"))
                          .values())
        if any(PLACEHOLDER_PATTERN.match(t) for t in texts):
            tbd.append(get_idshort(cd))
    rows.append(_row("CD 정의 작성 필요 ([작성 필요]/[TBD])", tbd, warn=True))
    parent_map = build_parent_map(root)
    samples = [
        element_path(p, parent_map) for p in find_all_by_name(root, "property")
        if element_path(p, parent_map) in AUTO_VALUES
        or text_of(find_direct_child(p, "value")) in SAMPLE_LIKE
    ]
    rows.append(_row("예시값 변경 필요 (임의 입력)", samples, warn=True))

    # ---- IDTA 2-14 ----
    empties = []
    for element in root.iter():
        name = local_name(element.tag)
        if len(list(element)) == 0 and element.text is not None \
                and not element.text.strip():
            empties.append(name)
        elif len(list(element)) == 0 and element.text is None and name in (
            OPTIONAL_STRINGS | {"value", "text", "language", "id", "idShort"}
        ) and local_name(element.tag) != "orderRelevant":
            empties.append(name)
    rows.append(_row("IDTA 2-14 문자열 최소 길이", empties))

    # ---- 참조 무결성 ----
    sm_ids = {get_id(sm) for sm in submodels}
    dangling = []
    for reference in root.iter():
        keys = find_direct_child(reference, "keys")
        if keys is None or text_of(
            find_direct_child(reference, "type")
        ) != "ModelReference":
            continue
        key_list = [k for k in keys if local_name(k.tag) == "key"]
        if key_list and text_of(find_direct_child(key_list[0], "type")) == \
                "Submodel":
            value = text_of(find_direct_child(key_list[0], "value"))
            if value not in sm_ids and not is_standard_id(value) and \
                    local_name(reference.tag) not in SEMANTIC_REFERENCE_TAGS:
                dangling.append(value)
    rows.append(_row("참조 무결성 (ModelReference)", dangling))

    # ---- IDTA 3-x / 4-x (aas-core) ----
    core = verify_with_aas_core(xml_bytes)
    if core is None:
        rows.append({
            "검사항목": "IDTA 3-x/4-x 메타모델 제약",
            "결과": "➖ 미실행",
            "건수": 0,
            "상세": "aas-core3.0 / aas-core3.1 설치 시 검사",
        })
    else:
        template = [p for p in core if "concept_descriptions" in p["경로"]]
        other = [p for p in core if p not in template]
        rows.append(_row(
            "IDTA 3-x/4-x 메타모델 제약 (자체 요소)",
            [f"{p['경로']}: {p['오류']}" for p in other],
        ))
        rows.append(_row(
            "IDTA 3-x/4-x 메타모델 제약 (CD, 템플릿 원본 포함)",
            [f"{p['경로']}: {p['오류']}" for p in template],
            warn=True,
        ))
    return rows


def verify_with_aas_core(xml_bytes):
    """네임스페이스(3.0/3.1)에 맞는 aas-core 로 메타모델 제약을 검사한다."""
    namespace = get_namespace(ET.fromstring(xml_bytes).tag)
    try:
        if namespace == NS_AAS31:
            import aas_core3_1.xmlization as xmlization
            import aas_core3_1.verification as verification
        else:
            import aas_core3.xmlization as xmlization
            import aas_core3.verification as verification
    except ImportError:
        return None

    try:
        environment = xmlization.environment_from_str(
            xml_bytes.decode("utf-8")
        )
    except Exception as error:  # noqa: BLE001
        return [{"경로": "(역직렬화 실패: 스키마 구조 오류)", "오류": str(error)}]

    problems = []
    for issue in verification.verify(environment):
        problems.append({"경로": str(issue.path), "오류": issue.cause})
        if len(problems) >= 300:
            break
    return problems


# 로그 '구분' -> Validator 항목 (화면/CSV 표시용)
LOG_ITEM_RULES = [
    ("idShort", "KOSMO-*-#1 idShort"),
    ("AAS ID", "KOSMO-AAS #3"),
    ("globalAssetId", "KOSMO-AAS #5"),
    ("assetKind", "KOSMO-AAS #6"),
    ("썸네일", "KOSMO-AAS #1"),
    ("defaultThumbnail", "KOSMO-AAS #1"),
    ("Submodel ID", "KOSMO-Submodel #2"),
    ("Submodel SemanticID", "KOSMO-Submodel #3"),
    ("Submodel kind", "KOSMO-Submodel #4"),
    ("특수 SME", "KOSMO-SME #2"),
    ("SemanticID", "KOSMO-SME #3"),
    ("CD 생성", "KOSMO-SME #3"),
    ("CD 신규", "KOSMO-SME #3"),
    ("Property value", "KOSMO-SME #4"),
    ("예시값", "예시값 (D-12)"),
    ("valueType", "KOSMO-SME #5"),
    ("dataType", "KOSMO-SME #5 / CD #4"),
    ("ConceptDescription ID", "KOSMO-CD #2"),
    ("CD 텍스트", "KOSMO-CD #3"),
    ("CD 정의 생성", "KOSMO-CD #3"),
    ("CD 한글 생성", "KOSMO-CD #3"),
    ("CD Description", "KOSMO-CD #3"),
    ("CD 한글", "KOSMO-CD #3"),
    ("CD 자리표시자", "KOSMO-CD #3"),
    ("Duplicate CD", "KOSMO-CD #5"),
    ("빈 ", "IDTA 2-14"),
    ("AASc-3a-009", "IDTA 4-7 (AASc-3a-009)"),
    ("AASd-107", "IDTA 3-5 (AASd-107)"),
    ("AASd-108", "IDTA 3-6 (AASd-108)"),
    ("AASd-109", "IDTA 3-7 (AASd-109)"),
    ("SML", "IDTA 3-5~3-7 (SML)"),
    ("요소 순서", "IDTA 2-x (XSD)"),
    ("D-05", "패키지 (D-05)"),
    ("Content_Types", "패키지 (D-05)"),
    ("ZIP", "패키지 (D-05)"),
    ("참조", "참조 무결성"),
]


def annotate_logs(logs):
    for log in logs:
        category = log.get("구분", "")
        log["검사항목"] = next(
            (item for key, item in LOG_ITEM_RULES if key in category), "기타"
        )
    return logs


# =========================================================
# 한글/영문 사전 CSV
# =========================================================

GLOSSARY_COLUMNS = {
    "idshort": "idShort",
    "ko_name": "ko_name", "한글명": "ko_name", "preferredname_ko": "ko_name",
    "ko_definition": "ko_definition", "한글정의": "ko_definition",
    "definition_ko": "ko_definition",
    "en_name": "en_name", "영문명": "en_name", "preferredname_en": "en_name",
    "en_definition": "en_definition", "영문정의": "en_definition",
    "definition_en": "en_definition",
}

GLOSSARY_TEMPLATE_CSV = (
    "idShort,ko_name,ko_definition,en_name,en_definition\n"
    "WindingSpeedSetpoint,권취 속도 설정값,권취축 회전 속도의 설정값,"
    "Winding Speed Setpoint,Setpoint of the rotational speed of the "
    "winding shaft\n"
)


def parse_glossary_csv(data):
    """CSV(UTF-8/UTF-8-BOM/CP949) -> {idShort: {...}}"""
    import csv
    text = None
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 인코딩을 읽을 수 없습니다 (UTF-8 또는 CP949).")
    reader = csv.DictReader(io.StringIO(text))
    mapping = {}
    for raw in reader:
        row = {}
        for key, value in raw.items():
            if key is None:
                continue
            normalized = GLOSSARY_COLUMNS.get(
                key.strip().lower().replace(" ", "")
            )
            if normalized and value and value.strip():
                row[normalized] = value.strip()
        idshort = row.pop("idShort", None)
        # 자리표시자가 그대로 남은 칸은 무시
        row = {k: v for k, v in row.items() if not PLACEHOLDER_PATTERN.match(v)}
        if idshort:
            mapping[sanitize_idshort(idshort) or idshort] = row
    return mapping


def export_value_review_csv(xml_bytes):
    """
    Property / Range 예시값 검토용 CSV.
    임의로 채운 값과 임의값으로 보이는 값을 위에 둔다.
    '새 값' 열을 채워 다시 올리면 반영된다.
    """
    import csv
    root = ET.fromstring(xml_bytes)
    parent_map = build_parent_map(root)
    rows = []
    for element in root.iter():
        name = local_name(element.tag)
        if name not in ("property", "range"):
            continue
        path = element_path(element, parent_map)
        if name == "property":
            current = text_of(find_direct_child(element, "value"))
        else:
            current = (
                f"{text_of(find_direct_child(element, 'min'))}.."
                f"{text_of(find_direct_child(element, 'max'))}"
            )
        if path in AUTO_VALUES:
            note = "이번 실행에서 임의 입력"
        elif current in SAMPLE_LIKE:
            note = "임의값으로 보임"
        else:
            note = ""
        rows.append({
            "경로": path,
            "종류": SME_TYPE_NAME[name],
            "valueType": text_of(find_direct_child(element, "valueType")),
            "현재값": current,
            "새 값": "",
            "검토": "필요" if note else "",
            "비고": note + (" / IDTA 템플릿 요소" if is_protected_element(
                element
            ) else ""),
        })
    rows.sort(key=lambda r: (r["검토"] != "필요",))
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=list(rows[0]) if rows else [
        "경로", "종류", "valueType", "현재값", "새 값", "검토", "비고",
    ])
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def parse_value_csv(data):
    """예시값 CSV -> {경로: 새 값} (새 값이 있는 행만)."""
    import csv
    text = None
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("CSV 인코딩을 읽을 수 없습니다 (UTF-8 또는 CP949).")
    mapping = {}
    for row in csv.DictReader(io.StringIO(text)):
        path = (row.get("경로") or "").strip()
        new_value = (row.get("새 값") or "").strip()
        if path and new_value:
            mapping[path] = new_value
    return mapping


def apply_value_overrides(root, overrides, logs):
    """CSV '새 값' 을 경로로 찾아 반영. 형식이 안 맞으면 반영하지 않고 경고."""
    if not overrides:
        return
    parent_map = build_parent_map(root)
    found = set()
    for element in root.iter():
        name = local_name(element.tag)
        if name not in ("property", "range"):
            continue
        path = element_path(element, parent_map)
        if path not in overrides:
            continue
        found.add(path)
        new_value = overrides[path]
        value_type = text_of(find_direct_child(element, "valueType"))
        if name == "range":
            parts = re.split(r"\s*(?:\.\.|~)\s*", new_value)
            if len(parts) != 2 or not all(
                value_matches_type(p, value_type) for p in parts
            ):
                logs.append({
                    "구분": "예시값 반영 실패",
                    "요소": path,
                    "수정 전": new_value,
                    "수정 후": f"Range 는 '최소..최대' 형식, {value_type}",
                    "결과": "Warning",
                })
                continue
            for tag, part in zip(("min", "max"), parts):
                target = find_direct_child(element, tag)
                if target is None:
                    target = ET.Element(ns_tag(element, tag))
                    insert_ordered(element, target)
                target.text = part
        else:
            if not value_matches_type(new_value, value_type):
                logs.append({
                    "구분": "예시값 반영 실패",
                    "요소": path,
                    "수정 전": new_value,
                    "수정 후": f"valueType {value_type} 형식이 아님",
                    "결과": "Warning",
                })
                continue
            target = find_direct_child(element, "value")
            if target is None:
                target = ET.Element(ns_tag(element, "value"))
                insert_ordered(element, target)
            target.text = normalize_value_text(new_value, value_type)
        AUTO_VALUES.discard(path)
        logs.append({
            "구분": "예시값 반영",
            "요소": path,
            "수정 전": "",
            "수정 후": new_value,
            "결과": "Auto Fixed",
        })
    for path in sorted(set(overrides) - found):
        logs.append({
            "구분": "예시값 반영 실패",
            "요소": path,
            "수정 전": overrides[path],
            "수정 후": "해당 경로의 요소 없음",
            "결과": "Warning",
        })


def logs_to_csv(logs):
    import csv
    buffer = io.StringIO()
    fields = ["검사항목", "구분", "요소", "수정 전", "수정 후", "결과"]
    writer = csv.DictWriter(buffer, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for log in logs:
        writer.writerow(log)
    return buffer.getvalue().encode("utf-8-sig")



# =========================================================
# 검토 엑셀 (CD 이름·정의 + 예시값) 내보내기 / 읽기
# =========================================================

REVIEW_SHEET_GUIDE = "사용법"
REVIEW_SHEET_CD = "CD 목록"
REVIEW_SHEET_VALUE = "예시값"

CD_HEADERS = [
    "No", "구분", "요소 종류", "idShort", "사용 위치", "영문명", "한글명",
    "영문 정의", "한글 정의", "상태", "사용 요소 수", "CD id",
]
VALUE_HEADERS = [
    "No", "경로", "idShort", "valueType", "현재값", "새 값", "비고", "상태",
]


def _clean_text(text):
    """자리표시자는 빈칸으로 내보낸다 (엑셀에서 누락으로 보이게)."""
    return "" if not text or PLACEHOLDER_PATTERN.match(text) else text


def collect_review_rows(xml_bytes):
    """엑셀/화면 공통: CD 행과 예시값 행을 만든다."""
    root = ET.fromstring(xml_bytes)
    parent_map = build_parent_map(root)
    cd_by_id = {get_id(cd): cd for cd in get_concept_descriptions(root)}

    usage = {}          # cd id -> [element, ...] (문서 순서)
    for element in root.iter():
        if local_name(element.tag) not in SME_TYPES:
            continue
        semantic = get_semantic_id_value(element)
        if semantic:
            usage.setdefault(semantic, []).append(element)

    cd_rows = []
    for cd_id, elements in usage.items():
        cd = cd_by_id.get(cd_id)
        if cd is None:
            continue
        content = get_iec61360_content(cd)
        preferred = lang_map(find_direct_child(content, "preferredName")) \
            if content is not None else {}
        definition = lang_map(find_direct_child(content, "definition")) \
            if content is not None else {}
        first = elements[0]
        kinds = list(dict.fromkeys(
            SME_TYPE_NAME.get(local_name(e.tag), local_name(e.tag))
            for e in elements
        ))
        row = {
            "구분": "표준(IDTA·IRDI)" if is_standard_id(cd_id) else "자체",
            "요소 종류": ", ".join(kinds),
            "idShort": get_idshort(cd),
            "사용 위치": element_path(first, parent_map),
            "영문명": _clean_text(preferred.get("en")),
            "한글명": _clean_text(preferred.get("ko")),
            "영문 정의": _clean_text(definition.get("en")),
            "한글 정의": _clean_text(definition.get("ko")),
            "사용 요소 수": len(elements),
            "CD id": cd_id,
        }
        row["누락"] = [
            key for key in ("영문명", "한글명", "영문 정의", "한글 정의")
            if not row[key]
        ]
        cd_rows.append(row)

    value_rows = []
    for element in root.iter():
        name = local_name(element.tag)
        if name not in ("property", "range"):
            continue
        path = element_path(element, parent_map)
        if name == "property":
            current = text_of(find_direct_child(element, "value"))
        else:
            current = (
                f"{text_of(find_direct_child(element, 'min'))}.."
                f"{text_of(find_direct_child(element, 'max'))}"
            )
        if path in AUTO_VALUES:
            note = "임의 입력 - 변경 필요"
        elif current in SAMPLE_LIKE:
            note = "임의값으로 보임 - 확인 필요"
        else:
            note = ""
        value_rows.append({
            "경로": path,
            "idShort": get_idshort(element) or path.rsplit("/", 2)[-2],
            "valueType": text_of(find_direct_child(element, "valueType")),
            "현재값": current,
            "비고": note,
        })
    return cd_rows, value_rows


def review_counts(xml_bytes):
    cd_rows, value_rows = collect_review_rows(xml_bytes)
    return {
        "cd_total": len(cd_rows),
        "cd_missing": sum(1 for r in cd_rows if r["누락"]),
        "value_total": len(value_rows),
        "value_todo": sum(1 for r in value_rows if r["비고"]),
    }


def export_review_workbook(xml_bytes, file_label="AASX"):
    """사용법 / CD 목록 / 예시값 3개 시트의 검토 엑셀을 만든다."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    cd_rows, value_rows = collect_review_rows(xml_bytes)

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title = Font(name="Arial", size=14, bold=True)
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    input_fill = PatternFill("solid", fgColor="FFF2CC")     # 채울 칸
    missing_fill = PatternFill("solid", fgColor="FFD966")   # 비어 있는 칸
    locked_fill = PatternFill("solid", fgColor="EDEDED")    # 수정 안 함
    done_font = Font(name="Arial", size=10, color="2E7D32", bold=True)
    todo_font = Font(name="Arial", size=10, color="C00000", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()

    # ---------------- 사용법 ----------------
    guide = wb.active
    guide.title = REVIEW_SHEET_GUIDE
    n_cd = len(cd_rows) + 1
    n_val = len(value_rows) + 1
    lines = [
        ("AAS 검토 엑셀", title),
        (f"대상: {file_label}", font),
        ("", font),
        ("진행 순서", bold),
        ("1. 'CD 목록' 시트에서 상태가 '누락'인 행의 노란 빈칸(영문명·한글명·"
         "영문 정의·한글 정의)을 채웁니다.", font),
        ("2. '예시값' 시트에서 상태가 '변경 필요'인 행의 '새 값'을 채웁니다. "
         "Range 는 '최소..최대' 형식입니다.", font),
        ("3. 이 파일을 엑셀 형식(.xlsx) 그대로 저장합니다. 시트 이름과 머리글은 "
         "바꾸지 마세요.", font),
        ("4. Auto Fixer 화면에서 '원본 AASX'와 이 엑셀을 함께 올리고 다시 "
         "실행합니다.", font),
        ("5. 다시 받은 검토 엑셀의 누락이 0건이 될 때까지 반복합니다.", font),
        ("", font),
        ("칸 색깔", bold),
        ("진한 노랑: 비어 있어 채워야 하는 칸 (채우면 색이 빠집니다)", font),
        ("연한 노랑: 수정할 수 있는 칸", font),
        ("회색: 바꿀 수 없는 칸 (IDTA·표준 CD 의 영문 원본, 식별 정보). "
         "바꿔도 반영되지 않습니다.", font),
        ("", font),
        ("현재 남은 작업", bold),
    ]
    for index, (text, style) in enumerate(lines, start=1):
        cell = guide.cell(row=index, column=1, value=text)
        cell.font = style
    row = len(lines) + 1
    summary = [
        ("CD 누락",
         f"=COUNTIF('{REVIEW_SHEET_CD}'!J2:J{n_cd},\"누락\")",
         f"=COUNTA('{REVIEW_SHEET_CD}'!D2:D{n_cd})"),
        ("예시값 변경 필요",
         f"=COUNTIF('{REVIEW_SHEET_VALUE}'!H2:H{n_val},\"변경 필요\")",
         f"=COUNTA('{REVIEW_SHEET_VALUE}'!B2:B{n_val})"),
    ]
    guide.cell(row=row, column=1, value="항목").font = bold
    guide.cell(row=row, column=2, value="남은 건수").font = bold
    guide.cell(row=row, column=3, value="전체").font = bold
    for offset, (label, todo, total) in enumerate(summary, start=1):
        guide.cell(row=row + offset, column=1, value=label).font = font
        guide.cell(row=row + offset, column=2, value=todo).font = todo_font
        guide.cell(row=row + offset, column=3, value=total).font = font
    example_row = row + len(summary) + 2
    guide.cell(row=example_row, column=1, value="작성 예시").font = bold
    example = [
        ("idShort", "영문명", "한글명", "영문 정의", "한글 정의"),
        ("WindingSpeedSetpoint", "Winding Speed Setpoint", "권취 속도 설정값",
         "Setpoint of the rotational speed of the winding shaft",
         "권취축 회전 속도의 설정값"),
    ]
    for r_offset, values in enumerate(example, start=1):
        for c_offset, value in enumerate(values, start=1):
            cell = guide.cell(row=example_row + r_offset, column=c_offset,
                              value=value)
            cell.font = bold if r_offset == 1 else font
            cell.border = border
    guide.column_dimensions["A"].width = 30
    for col in "BCDE":
        guide.column_dimensions[col].width = 34

    # ---------------- CD 목록 ----------------
    ws = wb.create_sheet(REVIEW_SHEET_CD)
    for col, header in enumerate(CD_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for index, data in enumerate(cd_rows, start=2):
        standard = data["구분"].startswith("표준")
        values = [
            index - 1, data["구분"], data["요소 종류"], data["idShort"],
            data["사용 위치"], data["영문명"], data["한글명"],
            data["영문 정의"], data["한글 정의"],
            f'=IF(COUNTBLANK(F{index}:I{index})=0,"완료","누락")',
            data["사용 요소 수"], data["CD id"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            cell.font, cell.border, cell.alignment = font, border, wrap
            if col in (6, 7, 8, 9):
                # 표준 CD 영문은 원문이 있을 때만 잠금 (비어 있으면 입력 가능)
                locked = standard and col in (6, 8) and bool(value)
                cell.fill = locked_fill if locked else input_fill
                cell.number_format = "@"
            elif col != 10:
                cell.fill = locked_fill
    last = len(cd_rows) + 1
    if cd_rows:
        # 비어 있는 입력칸은 진하게 (채우면 바로 사라짐)
        ws.conditional_formatting.add(
            f"F2:I{last}",
            FormulaRule(formula=["LEN(TRIM(F2))=0"], fill=missing_fill),
        )
        ws.conditional_formatting.add(
            f"J2:J{last}",
            FormulaRule(formula=['J2="누락"'], font=todo_font),
        )
        ws.conditional_formatting.add(
            f"J2:J{last}",
            FormulaRule(formula=['J2="완료"'], font=done_font),
        )
    widths = [6, 14, 24, 28, 46, 28, 22, 50, 40, 9, 11, 46]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:L{last}"

    # ---------------- 예시값 ----------------
    vs = wb.create_sheet(REVIEW_SHEET_VALUE)
    for col, header in enumerate(VALUE_HEADERS, start=1):
        cell = vs.cell(row=1, column=col, value=header)
        cell.font, cell.fill, cell.border = header_font, header_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ordered = sorted(value_rows, key=lambda r: r["비고"] == "")
    for index, data in enumerate(ordered, start=2):
        values = [
            index - 1, data["경로"], data["idShort"], data["valueType"],
            data["현재값"], "", data["비고"],
            f'=IF(G{index}="","",IF(LEN(TRIM(F{index}))=0,"변경 필요",'
            f'"입력됨"))',
        ]
        for col, value in enumerate(values, start=1):
            cell = vs.cell(row=index, column=col, value=value)
            cell.font, cell.border, cell.alignment = font, border, wrap
            if col == 6:
                cell.fill = input_fill
                cell.number_format = "@"
            elif col == 5:
                cell.fill = locked_fill
                cell.number_format = "@"
            elif col != 8:
                cell.fill = locked_fill
    last_v = len(ordered) + 1
    if ordered:
        vs.conditional_formatting.add(
            f"F2:F{last_v}",
            FormulaRule(formula=['AND(G2<>"",LEN(TRIM(F2))=0)'],
                        fill=missing_fill),
        )
        vs.conditional_formatting.add(
            f"H2:H{last_v}",
            FormulaRule(formula=['H2="변경 필요"'], font=todo_font),
        )
        vs.conditional_formatting.add(
            f"H2:H{last_v}",
            FormulaRule(formula=['H2="입력됨"'], font=done_font),
        )
    for col, width in enumerate([6, 64, 28, 12, 26, 26, 26, 10], start=1):
        vs.column_dimensions[get_column_letter(col)].width = width
    vs.freeze_panes = "C2"
    vs.auto_filter.ref = f"A1:H{last_v}"

    # 인쇄: 가로, 폭 맞춤
    for sheet in (guide, ws, vs):
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    vs.print_title_rows = "1:1"

    wb.active = 1  # CD 목록 시트를 먼저 보여준다
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _cell_text(value):
    """엑셀이 숫자·날짜로 바꾼 값을 AAS 값 문자열로 되돌린다."""
    import datetime
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime.datetime):
        return value.isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_review_workbook(data):
    """
    검토 엑셀 -> (glossary_by_id, value_overrides)
    시트 이름과 머리글로 찾으므로 열 순서가 바뀌어도 읽는다.
    """
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    glossary_by_id, overrides = {}, {}

    def rows_of(sheet_name):
        if sheet_name not in wb.sheetnames:
            return []
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        return [dict(zip(headers, r)) for r in rows[1:]]

    key_map = {
        "영문명": "en_name", "한글명": "ko_name",
        "영문 정의": "en_definition", "한글 정의": "ko_definition",
    }
    for row in rows_of(REVIEW_SHEET_CD):
        cd_id = _cell_text(row.get("CD id"))
        if not cd_id:
            continue
        entry = {}
        for header, key in key_map.items():
            text = _cell_text(row.get(header))
            if text and not PLACEHOLDER_PATTERN.match(text):
                entry[key] = text
        if entry:
            glossary_by_id[cd_id] = entry

    for row in rows_of(REVIEW_SHEET_VALUE):
        path = _cell_text(row.get("경로"))
        new_value = _cell_text(row.get("새 값"))
        if path and new_value:
            overrides[path] = new_value

    if REVIEW_SHEET_CD not in wb.sheetnames and \
            REVIEW_SHEET_VALUE not in wb.sheetnames:
        raise ValueError(
            f"'{REVIEW_SHEET_CD}' / '{REVIEW_SHEET_VALUE}' 시트를 찾지 "
            "못했습니다. Auto Fixer 에서 받은 검토 엑셀인지 확인하세요."
        )
    return glossary_by_id, overrides
